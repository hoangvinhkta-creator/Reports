"""R3 §4 — XUẤT EXCEL từ EFFECTIVE DATA, không từ kết quả pipeline thô.

## Vì sao không dùng lại `excel_exporter.py`

`excel_exporter` xuất từ `ImportResult` — kết quả của MỘT lần chạy máy, đúng
tại thời điểm chạy. Nhưng từ PHB-03, thứ Owner nhìn trên màn hình KHÔNG phải
con số đó: giá nhập tay, việc gán lại nhân viên, phân loại Gia dụng cấp dòng
và việc loại một dòng khỏi báo cáo đều được hợp nhất LÚC ĐỌC
(`business_queries.build_lines`). R3 thêm loại dòng vào đúng chỗ ấy.

Nghĩa là một file xuất từ `ImportResult` là ảnh chụp của trạng thái TRƯỚC mọi
quyết định của Owner. Nó trông đầy đủ, nó cân, và nó nói một bộ số khác bộ số
trên màn hình — loại sai nguy hiểm nhất, vì không có gì trong file báo rằng
nó cũ.

Module này vì thế nhận đúng thứ màn hình đang hiển thị: các `BusinessLine` đã
hợp nhất, cộng phần chi tiết đi kèm. Nó KHÔNG tính lại một chỉ tiêu nào —
`totals()` ở đây là `business_metrics.totals`, đúng hàm mà trang Báo cáo gọi.

## Một cột "Giá nhập KPI", và một cột nói nó từ đâu

Yêu cầu R3: MỘT cột giá nhập KPI. Không có cột "giá MIN" cạnh cột "giá tay"
cạnh cột "giá chính sách" — ba cột như thế buộc người đọc tự chọn cột nào
đúng, và họ sẽ chọn khác nhau.

Con số trong cột đó là `BusinessLine.purchase_price` — giá nhập HIỆU LỰC, một
thẩm quyền đã sắp thứ tự (giá tay → giá tự động → chính sách loại dòng). Cột
`Nguồn giá nhập` đi kèm nói nó đến từ đâu, và cột `Giá tự động lúc sửa` giữ
provenance mà `R2 §4.4` yêu cầu: khi Owner thay một giá AUTO, con số bị thay
vẫn đọc lại được.

Ô TRỐNG nghĩa là CHƯA CÓ GIÁ. Không bao giờ là `0` (`OD-105B-01` §3).
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.reporting import business_metrics as bm
from app.modules.reporting import line_type as line_type_module
from app.modules.reporting import profit_gate

MONEY_FORMAT = '#,##0;[Red](#,##0);"—"'
RATE_FORMAT = '0.00%'
DATE_FORMAT = "dd/mm/yyyy"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)

#: Cột của bảng kê chi tiết. MỘT cột giá nhập KPI (`§4`), và các cột
#: provenance đi kèm nó — không phải thay nó.
LINE_COLUMNS = (
    ("Ngày bán", 12),
    ("Số BH", 14),
    ("Mặt hàng", 42),
    ("Loại dòng", 18),
    ("SL", 8),
    ("Đơn giá bán", 14),
    ("Chiết khấu", 13),
    ("Doanh thu", 15),
    ("Giá nhập KPI", 15),
    ("Nguồn giá nhập", 18),
    ("Giá tự động lúc sửa", 18),
    ("Người nhập giá", 16),
    ("Lý do sửa giá", 30),
    ("Lợi nhuận KPI", 15),
    ("Tỉ lệ quy đổi", 13),
    ("DS quy đổi", 15),
    ("Nhân viên", 20),
    ("Nguồn nhân viên", 15),
    ("Chưa chốt được vì", 42),
    ("Ngoại lệ gắn dòng", 34),
)

SUMMARY_COLUMNS = (("Chỉ tiêu", 34), ("Giá trị", 20), ("Ghi chú", 56))

#: Nhãn provenance viết cho người đọc, không cho lập trình viên.
PROVENANCE_LABELS = {
    bm.PROVENANCE_AUTO: "Tự động (MIN theo ngày bán)",
    bm.PROVENANCE_MANUAL: "Owner nhập tay",
    bm.PROVENANCE_MANUAL_OVERRIDE: "Owner sửa (thay giá tự động)",
    bm.PROVENANCE_POLICY_ZERO: "Chính sách (dòng phụ = 0)",
    bm.PROVENANCE_PENDING: "Chưa có giá",
}


class ExportIntegrityError(ValueError):
    """File xuất ra sẽ KHÔNG khớp giao diện — dừng, không ghi ra đĩa.

    Một file lệch mà vẫn ghi được là một file sẽ được gửi đi và được tin.
    """


@dataclass(frozen=True)
class ExportSheet:
    """Một sheet của file xuất: nhãn + đúng tập dòng đã hiển thị."""

    key: str
    label: str
    details: Sequence[dict]

    @property
    def lines(self) -> list:
        return [detail["line"] for detail in self.details]

    @property
    def totals(self) -> bm.BusinessTotals:
        return bm.totals(self.lines)


def build_workbook(
    *, period_label: str, sheets: Sequence[ExportSheet],
    period_totals: bm.BusinessTotals, generated_at: Optional[datetime] = None,
    closed=None, binding_exceptions: Optional[dict] = None,
    discount_double_count: Sequence[str] = (),
    excluded: Sequence[dict] = (),
) -> Workbook:
    """Dựng workbook. `period_totals` là chỉ tiêu MÀN HÌNH đang hiện.

    Bất biến được kiểm ngay tại đây, trước khi có một byte nào ra đĩa: tổng
    các sheet phải đúng bằng tổng kỳ. `reporting_sheets.sheet_key_of` là một
    hàm toàn phần nên điều đó đúng theo cấu tạo — nhưng "theo cấu tạo" là một
    lập luận, còn phép kiểm là một bằng chứng, và đây là file rời khỏi hệ
    thống.
    """
    _assert_sheets_add_up(sheets, period_totals)

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_summary(
        workbook, period_label=period_label, sheets=sheets,
        period_totals=period_totals, generated_at=generated_at, closed=closed,
        discount_double_count=discount_double_count, excluded=excluded)
    _write_by_employee(workbook, sheets, period_totals)
    for sheet in sheets:
        _write_lines(workbook, sheet, binding_exceptions or {})
    return workbook


def write_workbook(path: Path, workbook: Workbook) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def _assert_sheets_add_up(sheets, period_totals) -> None:
    """`§42` — phân hoạch, nên các chỉ tiêu CỘNG ĐƯỢC phải cộng lại đúng."""
    total_lines = sum(len(sheet.details) for sheet in sheets)
    if total_lines != period_totals.lines:
        raise ExportIntegrityError(
            f"Các sheet cộng lại {total_lines} dòng nhưng kỳ có "
            f"{period_totals.lines} — file xuất sẽ không khớp giao diện.")
    for name, reader in (
        ("doanh thu", lambda t: t.sales_revenue),
        ("lợi nhuận KPI", lambda t: t.kpi_profit),
    ):
        parts = [reader(sheet.totals) for sheet in sheets]
        if _sum(parts) != reader(period_totals):
            raise ExportIntegrityError(
                f"Tổng {name} của các sheet khác tổng kỳ — file xuất sẽ không "
                "khớp giao diện.")


def _sum(values):
    kept = [value for value in values if value is not None]
    return sum(kept, Decimal(0)) if kept else None


def _write_summary(
    workbook, *, period_label, sheets, period_totals, generated_at, closed,
    discount_double_count, excluded,
) -> None:
    sheet = workbook.create_sheet("Tổng kỳ")
    _header(sheet, SUMMARY_COLUMNS)
    rows = [
        ("Kỳ báo cáo", period_label, ""),
        ("Xuất lúc", (generated_at or datetime.now()).strftime("%d/%m/%Y %H:%M"),
         "Giờ máy chủ."),
        ("Trạng thái bộ số", _state_label(period_totals),
         "CHÍNH THỨC chỉ khi mọi dòng của kỳ đều đã góp một con số lợi nhuận "
         "KPI (coverage 100 %)."),
        ("Chốt kỳ", _closed_label(closed),
         "Kỳ đã chốt: mọi thay đổi phải đi qua một lần MỞ LẠI có ghi lý do."),
        ("Số dòng", period_totals.lines, ""),
        ("Số đơn", period_totals.orders,
         "Đếm theo Số BH khác nhau. Cộng các sheet lại sẽ LỚN HƠN con số này "
         "khi một đơn có hai nhân viên — đó là sự thật, không phải lỗi."),
        ("Doanh thu bán hàng", period_totals.sales_revenue,
         "Σ(Đơn giá × SL − Chiết khấu)."),
        ("Tổng SP đủ điều kiện", period_totals.qualifying_quantity,
         "Chỉ dòng có đơn giá bán > 1.000.000 đ."),
        ("Lợi nhuận KPI", period_totals.kpi_profit,
         "Σ((Đơn giá − Giá nhập KPI) × SL − Chiết khấu) trên các dòng chốt được."),
        ("— trong đó đã gán nhân viên", period_totals.employee_attributed_profit, ""),
        ("— chưa rõ nhân viên", period_totals.unattributed_profit,
         "Vẫn nằm trong tổng của kỳ, chưa vào bảng của ai."),
        ("DS quy đổi", period_totals.converted_sales,
         "Chia theo TỪNG DÒNG rồi cộng — không bao giờ chia tổng cho một tỉ lệ "
         "trung bình."),
        ("Dòng đã chốt được lợi nhuận",
         f"{period_totals.coverage.covered_lines}/{period_totals.coverage.total_lines}", ""),
        ("Dòng chưa có giá nhập", period_totals.coverage.missing_price_lines,
         "Ô Giá nhập KPI trống nghĩa là CHƯA CÓ GIÁ — không bao giờ là 0."),
        ("Dòng chỉ còn thiếu giá nhập", period_totals.coverage.owner_fixable_lines,
         "Gõ một con số là dòng có lợi nhuận ngay."),
        ("Dòng bị loại khỏi báo cáo", len(excluded),
         "Owner đã loại; sổ kế toán gốc giữ nguyên."),
    ]
    for code, count in period_totals.coverage.blocked_lines:
        rows.append((f"Cửa chặn · {code}", count, profit_gate.label(code)))
    if discount_double_count:
        rows.append((
            "Cảnh báo trừ chiết khấu hai lần", len(discount_double_count),
            "Các đơn có CẢ dòng \"Chiết khấu\" riêng lẫn cột chiết khấu: "
            + ", ".join(discount_double_count)))
    rows.append(("", "", ""))
    rows.append(("SHEET", "Doanh thu", "Lợi nhuận KPI · DS quy đổi · số dòng"))
    for item in sheets:
        totals = item.totals
        rows.append((
            item.label, totals.sales_revenue,
            f"{_money(totals.kpi_profit)} · {_money(totals.converted_sales)} · "
            f"{totals.lines} dòng"))

    for values in rows:
        sheet.append(list(values))
        cell = sheet.cell(row=sheet.max_row, column=2)
        if isinstance(cell.value, Decimal):
            cell.number_format = MONEY_FORMAT
        cell.alignment = Alignment(horizontal="right" if isinstance(
            cell.value, (Decimal, int)) else "left")
    sheet.freeze_panes = "A2"


EMPLOYEE_COLUMNS = (
    ("Nhân viên", 24), ("Nhóm", 14), ("Số dòng", 10), ("Số đơn", 10),
    ("Doanh thu", 16), ("SP đủ điều kiện", 16), ("Lợi nhuận KPI", 16),
    ("DS quy đổi", 16), ("Đã chốt được", 14), ("Trạng thái", 16),
)


def _write_by_employee(workbook, sheets, period_totals) -> None:
    """Sheet "Theo nhân viên" — CHIỀU THỨ HAI của cùng tập dòng.

    Các sheet dòng ở dưới chia theo ĐƠN VỊ BÁO CÁO (Nội thành · Gia dụng · một
    nhân viên bán lẻ), vì đó là phân hoạch mà tỉ lệ quy đổi và Target sống
    trên. Nhưng câu hỏi "tháng này Vinh làm được bao nhiêu" không trả lời được
    từ một sheet gộp ba người.

    Đây là một PHÂN HOẠCH khác của CÙNG tập dòng, không phải một phép gộp thứ
    hai: `group_by_employee` chạy trên chính các `BusinessLine` đã hiển thị,
    nên nó cộng lại đúng bằng tổng kỳ. Cột `Số đơn` là ngoại lệ đã biết và
    được nói ra ở sheet Tổng kỳ — một đơn có hai nhân viên đếm ở cả hai dòng.

    Không có sheet DÒNG riêng cho từng nhân viên: hai bộ sheet cùng chứa cùng
    những dòng ấy sẽ làm mọi phép cộng trong file đếm đôi.
    """
    sheet = workbook.create_sheet("Theo nhân viên")
    _header(sheet, EMPLOYEE_COLUMNS)
    lines = [line for item in sheets for line in item.lines]
    for name, group, totals in bm.group_by_employee(lines):
        sheet.append([
            name or "Chưa xác định", group or "—", totals.lines, totals.orders,
            totals.sales_revenue, totals.qualifying_quantity, totals.kpi_profit,
            totals.converted_sales,
            f"{totals.coverage.covered_lines}/{totals.coverage.total_lines}",
            _state_label(totals),
        ])
        for column in (5, 7, 8):
            sheet.cell(row=sheet.max_row, column=column).number_format = MONEY_FORMAT
    sheet.append([
        "TỔNG KỲ", "", period_totals.lines, period_totals.orders,
        period_totals.sales_revenue, period_totals.qualifying_quantity,
        period_totals.kpi_profit, period_totals.converted_sales,
        f"{period_totals.coverage.covered_lines}/"
        f"{period_totals.coverage.total_lines}", _state_label(period_totals),
    ])
    for column in range(1, len(EMPLOYEE_COLUMNS) + 1):
        sheet.cell(row=sheet.max_row, column=column).font = TOTAL_FONT
    for column in (5, 7, 8):
        sheet.cell(row=sheet.max_row, column=column).number_format = MONEY_FORMAT
    sheet.freeze_panes = "A2"


def _write_lines(workbook, item: ExportSheet, binding_exceptions: dict) -> None:
    sheet = workbook.create_sheet(_safe_title(item.label))
    _header(sheet, LINE_COLUMNS)
    for detail in item.details:
        line = detail["line"]
        key = (detail["order_key"], detail["product_key"],
               detail["occurrence_index"])
        raised = binding_exceptions.get(key)
        sheet.append([
            detail["sale_date"],
            detail["order_key"],
            detail["product_raw"] or "",
            line_type_module.label(line.line_type),
            line.quantity,
            line.sell_price,
            line.discount,
            line.total_sales,
            # MỘT cột giá nhập KPI. Trống = chưa có giá, không phải 0.
            line.purchase_price,
            PROVENANCE_LABELS.get(line.purchase_provenance,
                                  line.purchase_provenance),
            detail.get("override_auto_price_at_entry"),
            detail.get("override_entered_by") or "",
            detail.get("override_reason") or "",
            line.kpi_profit,
            line.conversion_rate,
            line.converted_sales,
            line.employee or "Chưa xác định",
            "Owner gán" if line.employee_provenance == "MANUAL" else "Sổ ghi",
            " · ".join(profit_gate.label(code) for code in line.profit_blockers),
            _exception_note(raised),
        ])
        _format_line_row(sheet, sheet.max_row)

    totals = item.totals
    sheet.append([
        "TỔNG", "", f"{totals.lines} dòng · {totals.orders} đơn", "",
        totals.qualifying_quantity, None, None, totals.sales_revenue,
        None, "", None, "", "", totals.kpi_profit, None,
        totals.converted_sales, "", "", _state_label(totals), "",
    ])
    for column in range(1, len(LINE_COLUMNS) + 1):
        sheet.cell(row=sheet.max_row, column=column).font = TOTAL_FONT
    _format_line_row(sheet, sheet.max_row)
    sheet.freeze_panes = "A2"


def _format_line_row(sheet, row: int) -> None:
    for column in (6, 7, 8, 9, 11, 14, 16):
        sheet.cell(row=row, column=column).number_format = MONEY_FORMAT
    sheet.cell(row=row, column=1).number_format = DATE_FORMAT
    sheet.cell(row=row, column=15).number_format = RATE_FORMAT


def _header(sheet, columns) -> None:
    sheet.append([label for label, _width in columns])
    for index, (_label, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width


def _state_label(totals: bm.BusinessTotals) -> str:
    return ("CHÍNH THỨC" if totals.state == bm.STATE_OFFICIAL
            else "CHƯA HOÀN CHỈNH")


def _closed_label(closed) -> str:
    if closed is None:
        return "CHƯA CHỐT"
    who = f" bởi {closed.closed_by}" if closed.closed_by else ""
    return f"ĐÃ CHỐT lần {closed.version_no} lúc {closed.closed_at}{who}"


def _exception_note(raised) -> str:
    if raised is None:
        return ""
    decisions = ", ".join(raised.protected_decisions) or "một quyết định"
    return (f"Không ghép chắc chắn được với khoá cũ "
            f"{list(raised.candidate_occurrence_indexes)} — đang treo {decisions}")


def _money(value) -> str:
    return "—" if value is None else f"{value:,.0f}"


def _safe_title(label: str) -> str:
    """Tên sheet Excel: tối đa 31 ký tự, không chứa `[]:*?/\\`."""
    cleaned = "".join("-" if char in "[]:*?/\\" else char for char in label)
    return (cleaned or "Sheet")[:31]


__all__ = [
    "EMPLOYEE_COLUMNS", "ExportIntegrityError", "ExportSheet", "LINE_COLUMNS",
    "PROVENANCE_LABELS", "build_workbook", "write_workbook",
]
