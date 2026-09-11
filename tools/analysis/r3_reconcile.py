"""R3 — ĐỐI SOÁT trên dữ liệu thật khả dụng (hai kỳ golden đã ẩn danh).

    python3 -m tools.analysis.r3_reconcile

Script này KHÔNG phải một test và KHÔNG khẳng định con số nào là đúng. Nó ĐO,
trên đúng đường chạy production, bốn thứ mà R3 tuyên bố đã sửa — rồi in ra để
một người đọc và đối chiếu:

    §1  nạp lại CÙNG file      → không dòng nào bị coi là mới hay đã sửa
        nạp lại file ĐẢO THỨ TỰ → khoá dòng KHÔNG đổi, quyết định không trôi
    §2  loại dòng              → phân bố thật, và số dòng phí được chính sách
                                 cấp giá 0 theo `OD-105B-01` §3
    §3  effective data         → tổng kỳ == Σ(sheet) == Σ(nhân viên)
    §4  file xuất              → con số trong .xlsx == con số trên màn hình

Dữ liệu: `tests/fixtures/golden/period_2026_01.xlsx` và `period_2026_06.xlsx`
— hai kỳ nghiệp vụ THẬT của Tín Phát, đã ẩn danh theo `OD-GB-1`. Đây là dữ
liệu thật khả dụng duy nhất trong repo; sổ gốc không bao giờ được commit
(`DEC-108`).

Không có tham số, không ghi vào database thật: mỗi lần chạy dựng một SQLite
tạm rồi xoá. In ra là toàn bộ đầu ra.
"""

from __future__ import annotations

import io
import shutil
import tempfile
from collections import Counter
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine

import tools.db as history_db
from app.history import extraction
from app.composition import build_price_composition, run_import_production
from app.modules.exporting.excel_exporter import present_lines
from app.modules.importing.raw_reader import read_raw_rows
from app.modules.reporting import business_metrics as bm
from app.modules.reporting import line_type as lt
from app.modules.reporting import line_type_config
from app.web import business_service, business_store, history_store

REPO_ROOT = Path(__file__).resolve().parents[2]
CONFIG_DIR = REPO_ROOT / "config"
GOLDEN = REPO_ROOT / "tests" / "fixtures" / "golden"

PERIODS = (
    ("period_2026_01.xlsx", (2026, 1), date(2026, 1, 1), date(2026, 1, 31)),
    ("period_2026_06.xlsx", (2026, 6), date(2026, 6, 1), date(2026, 6, 30)),
)


def presented_for(path: Path):
    """Các dòng ĐÃ TRÌNH BÀY của một workbook — đúng thứ XLSX in ra.

    Đi qua `run_import_production` (seam production, `app/composition.py`) và
    `present_lines`, chứ không tự dựng lại: đó là đường mà bản Web thật đi, nên
    đo ở đây là đo đúng cái production làm. `run_import` trần sẽ thiếu bằng
    chứng giá và `present_lines` từ chối ngay — đúng như nó nên làm.
    """
    composition = build_price_composition(CONFIG_DIR)
    result = run_import_production(path, CONFIG_DIR, price_composition=composition)
    return result, present_lines(result, composition.records, read_raw_rows(path))


def build_service(engine):
    return business_service.BusinessReportService(
        engine=engine, store=business_store.BusinessDecisionStore(engine))


def write_snapshot(repository, presented, *, run_id, created_at, fingerprint,
                   source_lines=None):
    source = (extraction.build_source_lines(presented)
              if source_lines is None else source_lines)
    results = extraction.build_result_lines(presented, source)
    return repository.write_snapshot(
        run_id=run_id, created_at=created_at, source_file_name="so.xlsx",
        file_fingerprint=fingerprint, file_size=1,
        header_text=None, sheet_data_rows=len(source),
        rows_without_order_id=0, source_lines=source, result_lines=results,
        evidence={}, summary={})


def reversed_rows(presented):
    """Cùng các dòng, ĐẢO NGƯỢC thứ tự `source_row` — mô phỏng sổ bị sắp lại.

    Chỉ đảo VỊ TRÍ: mọi giá trị nghiệp vụ giữ nguyên tuyệt đối. Nếu khoá dòng
    phụ thuộc vị trí, đúng phép biến đổi này sẽ làm nó đổi.
    """
    rows = sorted({view.line.raw.source_row for view in presented})
    swapped = dict(zip(rows, reversed(rows)))
    source = extraction.build_source_lines(presented)
    moved = [
        type(item)(**{**{f: getattr(item, f) for f in item.__dataclass_fields__},
                      "source_row": swapped[item.source_row]})
        for item in source
    ]
    return sorted(moved, key=lambda item: item.source_row)


def report(title: str) -> None:
    print()
    print("=" * 78)
    print(title)
    print("=" * 78)


def main() -> None:
    vocabulary = line_type_config.load_vocabulary(CONFIG_DIR / "line_types.yaml")
    for filename, period, start, end in PERIODS:
        path = GOLDEN / filename
        report(f"{filename} — kỳ {period[1]:02d}/{period[0]}")
        _result, presented = presented_for(path)
        print(f"dòng đã trình bày: {len(presented)}")

        # --- §2 — loại dòng trên dữ liệu thật -------------------------
        types = Counter(
            lt.classify(
                vocabulary=vocabulary,
                order_key=view.line.order_id,
                product_raw=view.line.product_raw,
                sell_price=view.line.sell_price,
                quantity=view.line.quantity)
            for view in presented)
        print("\n§2 — phân bố loại dòng:")
        for name in lt.LINE_TYPES:
            if types.get(name):
                print(f"    {name:<20} {types[name]:>5}  ({lt.label(name)})")

        temp = Path(tempfile.mkdtemp())
        try:
            engine = create_engine(f"sqlite:///{temp / 'history.db'}")
            history_db.create_all_for_test(engine)
            repository = history_store.SnapshotRepository(engine)
            service = build_service(engine)

            first = write_snapshot(
                repository, presented, run_id="r1",
                created_at="2026-09-08T01:00:00", fingerprint="fp-1")
            print("\n§1 — nạp lần đầu:", dict(first.counts),
                  f"· ngoại lệ gắn dòng: {first.ambiguous_bindings}")

            same = write_snapshot(
                repository, presented, run_id="r2",
                created_at="2026-09-08T02:00:00", fingerprint="fp-1")
            print("§1 — nạp LẠI cùng file:", dict(same.counts),
                  f"· ngoại lệ gắn dòng: {same.ambiguous_bindings}")

            before = {
                (row.order_key, row.product_key, row.occurrence_index)
                for row in _current_keys(engine)}
            shuffled = write_snapshot(
                repository, presented, run_id="r3",
                created_at="2026-09-08T03:00:00", fingerprint="fp-2",
                source_lines=reversed_rows(presented))
            after = {
                (row.order_key, row.product_key, row.occurrence_index)
                for row in _current_keys(engine)}
            print("§1 — nạp lại file ĐẢO THỨ TỰ DÒNG:", dict(shuffled.counts),
                  f"· ngoại lệ gắn dòng: {shuffled.ambiguous_bindings}")
            print(f"       khoá dòng mới phát sinh: {len(after - before)}"
                  f" (0 = không quyết định nào có thể trôi)")

            # --- §3 — một effective data ------------------------------
            data = service.period(date_from=start, date_to=end, period=period)
            totals = data.totals
            by_employee = bm.group_by_employee(data.lines)
            by_sheet = [service_sheet_totals(service, data, sheet)
                        for sheet in service.sheets(data)]
            print("\n§3 — effective data:")
            print(f"    dòng kỳ                {totals.lines}")
            print(f"    Σ dòng theo nhân viên  "
                  f"{sum(t.lines for _n, _g, t in by_employee)}")
            print(f"    Σ dòng theo sheet      {sum(t.lines for t in by_sheet)}")
            print(f"    doanh thu kỳ           {_money(totals.sales_revenue)}")
            print(f"    Σ doanh thu nhân viên  "
                  f"{_money(_sum(t.sales_revenue for _n, _g, t in by_employee))}")
            print(f"    lợi nhuận KPI kỳ       {_money(totals.kpi_profit)}")
            print(f"    coverage               "
                  f"{totals.coverage.covered_lines}/{totals.coverage.total_lines}"
                  f"  ({totals.state})")
            print(f"    dòng chưa có giá nhập  "
                  f"{totals.coverage.missing_price_lines}")
            policy = sum(1 for line in data.lines
                         if line.purchase_provenance == bm.PROVENANCE_POLICY_ZERO)
            print(f"    dòng giá 0 theo chính sách (OD-105B-01 §3): {policy}")
            blocked = dict(totals.coverage.blocked_lines)
            print(f"    cửa chặn                {blocked}")
            print(f"    đơn nghi trừ chiết khấu hai lần: "
                  f"{list(data.discount_double_count)}")

            # --- §4 — file xuất nói cùng câu --------------------------
            stream = io.BytesIO()
            service.export_workbook(
                data=data, period_label=f"Tháng {period[1]:02d}/{period[0]}"
            ).save(stream)
            stream.seek(0)
            book = load_workbook(stream)
            exported = _exported_totals(book)
            print("\n§4 — file xuất:")
            print(f"    sheet                  {book.sheetnames}")
            print(f"    dòng trong file        {exported['lines']}"
                  f"  (màn hình: {totals.lines})")
            print(f"    lợi nhuận KPI trong file {_money(exported['kpi_profit'])}"
                  f"  (màn hình: {_money(totals.kpi_profit)})")
            print(f"    KHỚP                   "
                  f"{exported['lines'] == totals.lines and _eq(exported['kpi_profit'], totals.kpi_profit)}")
        finally:
            shutil.rmtree(temp, ignore_errors=True)


def service_sheet_totals(service, data, sheet):
    return data.for_sheet(sheet).totals


def _current_keys(engine):
    from tools.db.schema import order_line_current
    with engine.connect() as connection:
        return list(connection.execute(order_line_current.select()))


def _sum(values):
    kept = [value for value in values if value is not None]
    return sum(kept, Decimal(0)) if kept else None


def _eq(left, right) -> bool:
    if left is None or right is None:
        return left is right
    return Decimal(str(left)) == Decimal(str(right))


def _money(value) -> str:
    return "—" if value is None else f"{value:,.0f}"


def _exported_totals(book) -> dict:
    sheet = book["Tổng kỳ"]
    values = {}
    for row in range(2, sheet.max_row + 1):
        values[sheet.cell(row=row, column=1).value] = \
            sheet.cell(row=row, column=2).value
    return {"lines": values.get("Số dòng"),
            "kpi_profit": values.get("Lợi nhuận KPI")}


if __name__ == "__main__":
    main()
