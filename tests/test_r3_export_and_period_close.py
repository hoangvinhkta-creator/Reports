"""R3 §4 + §5 — file xuất đọc EFFECTIVE DATA, và chốt kỳ thật sự khoá.

Hai mệnh đề trung tâm:

    §4  Con số trong file .xlsx PHẢI đúng bằng con số trên màn hình. Đường
        xuất cũ (`excel_exporter` ← `ImportResult`) là ảnh chụp trạng thái
        TRƯỚC mọi quyết định của Owner — nó trông đầy đủ, nó cân, và nó nói
        một bộ số khác. Loại sai nguy hiểm nhất, vì file không tự khai là cũ.

    §5  Chốt kỳ phải TỪ CHỐI một lần ghi, không chỉ hiện một cảnh báo. Một
        cái chốt "khuyên nên cẩn thận" thì bộ số đã duyệt vẫn đổi được sau
        lưng người đã duyệt nó.
"""

from __future__ import annotations

import ast
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.modules.exporting import business_export
from app.modules.reporting import business_metrics as bm
from app.modules.reporting import line_type as lt
from app.web import period_lock

from tests.test_business_vertical import (  # noqa: F401 — fixtures
    JANUARY, engine, pair, persist, repository, service, store,
)

PERIOD = (2026, 1)


def workbook_of(service, data, *, label="Tháng 01/2026"):
    stream = io.BytesIO()
    service.export_workbook(data=data, period_label=label).save(stream)
    stream.seek(0)
    return load_workbook(stream)


def column_of(sheet, header: str) -> int:
    for index, cell in enumerate(sheet[1], start=1):
        if cell.value == header:
            return index
    raise AssertionError(f"Không có cột {header!r} trong sheet {sheet.title!r}")


def values(sheet, header: str) -> list:
    index = column_of(sheet, header)
    return [sheet.cell(row=row, column=index).value
            for row in range(2, sheet.max_row + 1)]


def summary_value(sheet, label: str):
    for row in range(2, sheet.max_row + 1):
        if sheet.cell(row=row, column=1).value == label:
            return sheet.cell(row=row, column=2).value
    raise AssertionError(f"Không có dòng {label!r} ở sheet Tổng kỳ")


# --- §4 — file xuất nói ĐÚNG câu mà màn hình nói --------------------------

class TestTheExportReadsWhatTheScreenReads:
    def test_it_carries_the_owner_price_not_the_pipeline_price(
        self, repository, service, store
    ):
        """Bài trung tâm của §4.

        Pipeline ghi 5.000.000. Owner sửa thành 4.000.000. File xuất PHẢI
        mang 4.000.000 — con số Owner nhìn thấy — chứ không phải con số của
        lần chạy máy.
        """
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        store.set_purchase_price(
            order_key="BH1", product_key=_product_key("Tủ lạnh Panasonic"),
            occurrence_index=1, price=Decimal("4000000"),
            auto_price=Decimal("5000000"), entered_by="owner-web",
            reason="Giá nhập thực tế theo hoá đơn")

        data = service.period(**JANUARY, period=PERIOD)
        book = workbook_of(service, data)
        sheet = book["Nội thành"]
        assert values(sheet, "Giá nhập KPI")[0] == 4000000
        assert values(sheet, "Nguồn giá nhập")[0] == \
            business_export.PROVENANCE_LABELS[bm.PROVENANCE_MANUAL_OVERRIDE]
        # Provenance mà `R2 §4.4` yêu cầu vẫn đọc lại được từ chính file.
        assert values(sheet, "Giá tự động lúc sửa")[0] == 5000000
        assert values(sheet, "Người nhập giá")[0] == "owner-web"
        assert values(sheet, "Lý do sửa giá")[0] == "Giá nhập thực tế theo hoá đơn"
        # Lợi nhuận đi theo giá mới: (8.000.000 − 4.000.000) × 1 − 0.
        assert values(sheet, "Lợi nhuận KPI")[0] == 4000000

    def test_there_is_exactly_one_kpi_purchase_price_column(self):
        headers = [label for label, _width in business_export.LINE_COLUMNS]
        assert headers.count("Giá nhập KPI") == 1
        # Không có cột "giá MIN" đứng cạnh: ba cột giá buộc người đọc tự chọn
        # cột nào đúng, và họ sẽ chọn khác nhau.
        assert not [h for h in headers if h != "Giá nhập KPI" and "Giá nhập" in h]

    def test_a_missing_price_is_an_empty_cell_never_zero(
        self, repository, service
    ):
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase=None,
            kpi_profit=None)])
        book = workbook_of(service, service.period(**JANUARY, period=PERIOD))
        assert values(book["Nội thành"], "Giá nhập KPI")[0] is None
        assert values(book["Nội thành"], "Lợi nhuận KPI")[0] is None

    def test_a_fee_line_shows_its_policy_zero_and_says_where_it_came_from(
        self, repository, service
    ):
        persist(repository, [pair(
            "BH1", product="Chi phí vận chuyển", sell="0", quantity="1",
            kpi_purchase=None, kpi_profit=None)])
        sheet = workbook_of(service, service.period(**JANUARY, period=PERIOD))["Nội thành"]
        assert values(sheet, "Giá nhập KPI")[0] == 0
        assert values(sheet, "Nguồn giá nhập")[0] == \
            business_export.PROVENANCE_LABELS[bm.PROVENANCE_POLICY_ZERO]
        assert values(sheet, "Loại dòng")[0] == lt.LINE_TYPE_LABELS[lt.TYPE_FEE]

    def test_the_totals_of_every_sheet_add_up_to_the_period(
        self, repository, service
    ):
        persist(repository, [
            pair("BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
                 kpi_profit="3000000", status="AUTO", row=6),
            pair("BH2", product="Máy giặt LG", employee="Ly", group="LE",
                 kpi_purchase="4000000", kpi_profit="2000000", status="AUTO",
                 row=7),
            pair("BH3", product="Bếp từ", employee=None, group=None,
                 kpi_purchase="1000000", kpi_profit="500000", status="AUTO",
                 row=8),
        ])
        data = service.period(**JANUARY, period=PERIOD)
        book = workbook_of(service, data)
        assert summary_value(book["Tổng kỳ"], "Số dòng") == data.totals.lines
        assert summary_value(book["Tổng kỳ"], "Lợi nhuận KPI") == \
            data.totals.kpi_profit
        # Mỗi dòng thuộc ĐÚNG một sheet — phân hoạch, không phải bộ lọc.
        line_sheets = [name for name in book.sheetnames
                       if name not in ("Tổng kỳ", "Theo nhân viên")]
        assert sum(book[name].max_row - 2 for name in line_sheets) == \
            data.totals.lines

    def test_a_sheet_that_does_not_add_up_is_refused_before_it_reaches_disk(self):
        """Bất biến `§42` được KIỂM, không chỉ được lập luận.

        "Đúng theo cấu tạo" là một lập luận; một phép kiểm là một bằng chứng,
        và đây là file rời khỏi hệ thống.
        """
        totals = bm.totals([])
        with pytest.raises(business_export.ExportIntegrityError):
            business_export.build_workbook(
                period_label="Tháng 01/2026",
                sheets=[business_export.ExportSheet(
                    key="x", label="X", details=[_detail()])],
                period_totals=totals)

    def test_an_excluded_line_is_absent_from_the_sheets_and_counted_separately(
        self, repository, service, store
    ):
        """`§30` — dòng Owner đã loại không góp vào chỉ tiêu nào, ở đâu cả."""
        persist(repository, [
            pair("BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
                 kpi_profit="3000000", status="AUTO", row=6),
            pair("BH2", product="Máy giặt LG", kpi_purchase="4000000",
                 kpi_profit="2000000", status="AUTO", row=7),
        ])
        store.exclude_line(order_key="BH2",
                           product_key=_product_key("Máy giặt LG"),
                           occurrence_index=1)
        data = service.period(**JANUARY, period=PERIOD)
        book = workbook_of(service, data)
        assert "BH2" not in values(book["Nội thành"], "Số BH")
        assert summary_value(book["Tổng kỳ"], "Dòng bị loại khỏi báo cáo") == 1

    def test_the_export_module_never_imports_the_pipeline_result(self):
        """Kiểm bằng CẤU TRÚC, đọc cây import — không quét văn bản.

        Quét văn bản sẽ vấp chính đoạn docstring giải thích vì sao module này
        KHÔNG đọc `ImportResult`, tức bắt đúng lời giải thích thay vì bắt lỗi.
        """
        tree = ast.parse(open(business_export.__file__, encoding="utf-8").read())
        imported = {node.module for node in ast.walk(tree)
                    if isinstance(node, ast.ImportFrom) and node.module}
        imported |= {alias.name for node in ast.walk(tree)
                     if isinstance(node, ast.Import) for alias in node.names}
        for forbidden in ("app.pipeline", "app.modules.exporting.excel_exporter",
                          "app.web.history_store", "tools.db.schema"):
            assert forbidden not in imported, forbidden
        # Và không tên nào của tầng pipeline xuất hiện trong MÃ (bỏ chú thích).
        names = {node.id for node in ast.walk(tree) if isinstance(node, ast.Name)}
        names |= {node.attr for node in ast.walk(tree)
                  if isinstance(node, ast.Attribute)}
        assert names.isdisjoint({
            "ImportResult", "run_import", "order_line_result_version"})


# --- §5 — chốt kỳ TỪ CHỐI, không chỉ cảnh báo -----------------------------

class TestClosingAPeriodActuallyLocksIt:
    def test_closing_then_writing_is_refused(self, repository, service, store):
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        data = service.period(**JANUARY, period=PERIOD)
        closed = service.close_period(period=PERIOD, data=data,
                                      closed_by="owner-web", note="Đã duyệt")
        assert closed.version_no == 1
        assert closed.line_count == 1

        with pytest.raises(period_lock.PeriodClosedError):
            service.guard_period_open(PERIOD)
        # Cửa chặn theo NGÀY BÁN của chính dòng, không theo kỳ đang xem.
        detail = service.period(**JANUARY, period=PERIOD).details[0]
        with pytest.raises(period_lock.PeriodClosedError):
            service.guard_line_open(detail)

    def test_another_period_stays_open(self, repository, service):
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        service.close_period(period=PERIOD,
                             data=service.period(**JANUARY, period=PERIOD))
        service.guard_period_open((2026, 2))  # không ném — kỳ khác vẫn mở

    def test_the_all_data_view_is_never_treated_as_closed(
        self, repository, service
    ):
        """Khung nhìn "toàn bộ dữ liệu" không phải một kỳ.

        Coi nó là đã chốt sẽ khoá MỌI thao tác ngay khi tháng đầu tiên được
        chốt — kể cả thao tác trên những tháng chưa ai duyệt.
        """
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        service.close_period(period=PERIOD,
                             data=service.period(**JANUARY, period=PERIOD))
        service.guard_period_open(None)

    def test_an_undated_line_is_locked_by_no_period(self, service):
        service.guard_line_open({"sale_date": None})

    def test_reopening_requires_a_reason_and_keeps_the_old_version(
        self, repository, service
    ):
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        data = service.period(**JANUARY, period=PERIOD)
        service.close_period(period=PERIOD, data=data, closed_by="owner-web")

        with pytest.raises(period_lock.MissingReopenReasonError):
            service.reopen_period(period=PERIOD, reason="   ")
        service.reopen_period(period=PERIOD, reason="Sót một hoá đơn",
                              reopened_by="owner-web")
        service.guard_period_open(PERIOD)  # đã mở lại — không còn ném

        # Lần chốt cũ KHÔNG bị xoá: nó chỉ được đánh dấu đã mở.
        history = service.period_store.history(year=2026, month=1)
        assert [item.version_no for item in history] == [1]
        # Chốt lại ⟹ một PHIÊN BẢN mới, không ghi đè bản cũ.
        service.close_period(period=PERIOD,
                             data=service.period(**JANUARY, period=PERIOD))
        assert [item.version_no for item in
                service.period_store.history(year=2026, month=1)] == [2, 1]

    def test_reopening_a_period_that_was_never_closed_is_refused(self, service):
        with pytest.raises(period_lock.PeriodNotClosedError):
            service.reopen_period(period=PERIOD, reason="không có gì để mở")

    def test_closing_twice_is_refused(self, repository, service):
        persist(repository, [pair("BH1", kpi_purchase="5000000",
                                  kpi_profit="3000000", status="AUTO")])
        data = service.period(**JANUARY, period=PERIOD)
        service.close_period(period=PERIOD, data=data)
        with pytest.raises(period_lock.PeriodClosedError):
            service.close_period(period=PERIOD, data=data)

    def test_the_snapshot_records_the_numbers_that_were_approved(
        self, repository, service
    ):
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        data = service.period(**JANUARY, period=PERIOD)
        service.close_period(period=PERIOD, data=data)
        closed = service.period_store.closed(year=2026, month=1)
        assert closed.totals["kpi_profit"] == str(data.totals.kpi_profit)
        assert closed.totals["state"] == data.totals.state

    def test_drift_is_reported_when_the_numbers_change_after_closing(
        self, repository, service, store
    ):
        """Một kỳ đã duyệt mà số đã đổi là một sự thật cần NÓI RA.

        Nó xảy ra thật: phân loại mã sản phẩm (thẩm quyền R2) là quyết định
        TOÀN CỤC và không đi qua cửa chặn kỳ, nên nó có thể đổi con số của một
        kỳ đã chốt. Cửa chặn không bắt được thì phép so vân tay phải bắt được.
        """
        persist(repository, [pair(
            "BH1", product="Tủ lạnh Panasonic", kpi_purchase="5000000",
            kpi_profit="3000000", status="AUTO")])
        data = service.period(**JANUARY, period=PERIOD)
        service.close_period(period=PERIOD, data=data)
        after = service.period(**JANUARY, period=PERIOD)
        assert service.period_drift(period=PERIOD, data=after) is False

        store.set_purchase_price(
            order_key="BH1", product_key=_product_key("Tủ lạnh Panasonic"),
            occurrence_index=1, price=Decimal("1000000"),
            auto_price=Decimal("5000000"), entered_by="ai-do",
            reason="ghi đè sau khi chốt")
        drifted = service.period(**JANUARY, period=PERIOD)
        assert service.period_drift(period=PERIOD, data=drifted) is True


def _product_key(product: str) -> str:
    from app.history import keys
    return keys.product_key(product)


def _detail() -> dict:
    line = bm.BusinessLine(
        order_key="BH1", employee="Vinh", employee_group="NOI_THANH",
        status="AUTO", sell_price=Decimal("1"), quantity=Decimal("1"),
        discount=Decimal("0"), total_sales=Decimal("1"),
        auto_purchase_price=Decimal("0"), auto_kpi_profit=Decimal("1"),
        kpi_authority_valid=True)
    return {"order_key": "BH1", "product_key": "pk", "occurrence_index": 1,
            "product_raw": "Tủ lạnh", "sale_date": date(2026, 1, 5),
            "line": line}
