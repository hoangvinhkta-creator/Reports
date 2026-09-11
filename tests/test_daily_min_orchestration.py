"""Điều phối R1: sổ bán hàng → kế hoạch hỏi giá → ảnh chụp MIN đúng KỲ.

Bộ này canh khoảng trống giữa "provider đọc đúng" và "hệ thống thật sự hỏi".
`tests/test_daily_min_contract.py` và `tests/test_daily_min_vertical.py` đều
NHẬN một capture đã có sẵn, nên cả hai xanh rực trong khi đường upload thật
không hề gọi hợp đồng `daily-min-v1` và mọi dòng Tracking Pending. Đó chính là
hình dạng lỗi đã lọt qua R1 vòng đầu.

Ba câu hỏi ở đây, và không câu nào trả lời được từ một capture cho sẵn:

1. Từ một workbook, hệ thống có suy ra ĐÚNG tập mã Tracking và khoảng ngày bán
   để hỏi không?
2. Khi trong kho có nhiều ảnh chụp, nó có chọn ảnh chụp phủ ĐÚNG kỳ đang chạy
   không — hay chọn "cái mới nhất" và ném cả kỳ ra ngoài cửa sổ?
3. Đường pull-on-run của bản web có thật sự gọi hợp đồng, đóng băng kết quả
   cho lần chạy ấy, rồi dọn file tạm không?
"""

from __future__ import annotations

import json
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest

from app import owner_usability
from app.modules.pricing.daily_min import load_daily_min_capture
from app.modules.pricing.daily_min.planning import (
    MAX_CONTRACT_DAYS,
    UnreadableSalesWorkbookError,
    plan_daily_min_request_for_workbook,
)
from app.modules.pricing.resolution.sources import load_tracking_catalog_capture
from app.modules.product.identity.store import JsonlProductIdentityStore
from tests.support import daily_min_fixtures as dmin
from tests.test_105e_price_composition import write_catalog_capture
from tests.test_daily_min_vertical import CATALOG, ROWS, write_sales
from tests.test_owner_usability import add_valid_captures
from tests.test_tracking_history_reader import CUTOVER
from tools.tracking import live_pull

AUG = date(2026, 8, 12)
SEP = date(2026, 9, 3)


def write_sales_over_days(path: Path, entries) -> Path:
    """Một workbook có nhiều NGÀY BÁN khác nhau.

    `write_sales` của bộ vertical ghi cả sổ vào một ngày, nên nó không dựng
    được cảnh "sổ trải rộng hơn một lượt gọi hợp đồng" — cảnh duy nhất chạm
    tới trần 62 ngày.
    """
    import openpyxl

    from tests.fixtures.synthetic_workbook import HEADER

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "SỔ CHI TIẾT BÁN HÀNG"
    sheet.append(["SỔ CHI TIẾT BÁN HÀNG"])
    sheet.append(["Từ ngày 01/01/2026 đến ngày 31/12/2026"])
    sheet.append([])
    sheet.append(HEADER)
    sheet.append(["", "", "Diễn giải chung"])
    for (order_id, product, quantity, sell), day in entries:
        sheet.append([
            day, order_id, f"Bán hàng {order_id}", product, f"KH{order_id}",
            f"Khách {order_id}", "1 Đường Test", "0900000000", quantity, sell,
            sell * quantity, 0, "Vũ Hạnh Ly 0868345633", "Shipper", 0, None, None,
        ])
        sheet.cell(sheet.max_row, 4).data_type = "s"
    book.save(path)
    book.close()
    return path


def catalog_snapshot(tmp_path: Path):
    return load_tracking_catalog_capture(write_catalog_capture(tmp_path, CATALOG))


def empty_store_view(tmp_path: Path):
    store = JsonlProductIdentityStore(log_path=tmp_path / "identity.jsonl")
    return store.read_at_revision(store.current_revision())


def plan_for(tmp_path: Path, sales: Path):
    return plan_daily_min_request_for_workbook(
        sales,
        tracking_catalog=catalog_snapshot(tmp_path),
        identity_store_view=empty_store_view(tmp_path),
        tracking_identity_authority=True,
    )


# ======================================================================
# 1. Kế hoạch hỏi giá — suy từ SỔ, không đoán, không hỏi thừa
# ======================================================================


def test_the_plan_asks_for_exactly_the_tracking_codes_the_ledger_uses(tmp_path):
    """Hỏi thừa thì chậm; hỏi thiếu thì một mã im lặng Pending."""
    plan = plan_for(tmp_path, write_sales(tmp_path / "s.xlsx", ROWS, day=SEP))
    assert plan is not None
    assert plan.product_codes == ("TRK-A", "TRK-B", "TRK-C", "TRK-D", "TRK-E")


def test_the_plan_spans_the_real_sale_dates_not_the_upload_day(tmp_path):
    """Khoảng ngày phải là khoảng NGÀY BÁN. Lấy ngày nạp sổ làm biên là đúng
    lỗi mà cả R1 tồn tại để chặn — chỉ dời một tầng lên trên."""
    rows = ROWS[:1]
    early = write_sales(tmp_path / "a.xlsx", rows, day=date(2026, 9, 3))
    late = write_sales(tmp_path / "b.xlsx", rows, day=date(2026, 9, 20))
    assert plan_for(tmp_path, early).date_from == date(2026, 9, 3)
    assert plan_for(tmp_path, late).date_to == date(2026, 9, 20)


def test_a_ledger_with_no_tracking_identity_has_nothing_to_ask(tmp_path):
    """`None` nghĩa là "không có câu hỏi nào", KHÔNG phải "mọi mã thiếu giá".

    Những dòng ấy đã Pending với lý do identity của chính chúng; hỏi Tracking
    bằng tên hàng trên chứng từ là đúng thứ tầng identity tồn tại để ngăn.
    """
    rows = [("BH9001", "Tên hàng chưa ai nhận diện", 1, 1_000_000)]
    assert plan_for(tmp_path, write_sales(tmp_path / "la.xlsx", rows, day=SEP)) is None


def test_a_ledger_with_no_rows_has_nothing_to_ask(tmp_path):
    assert plan_for(tmp_path, write_sales(tmp_path / "rong.xlsx", [], day=SEP)) is None


def test_a_plan_without_identity_sources_is_not_built_from_raw_names(tmp_path):
    """Thiếu danh mục ⇒ không có kế hoạch. Dựng kế hoạch từ `product_raw` thô
    là gửi tên hàng trên chứng từ sang Tracking như thể chúng là mã."""
    assert plan_daily_min_request_for_workbook(
        write_sales(tmp_path / "s.xlsx", ROWS, day=SEP),
        tracking_catalog=None,
        identity_store_view=empty_store_view(tmp_path),
    ) is None


def test_a_plan_knows_when_the_period_is_wider_than_one_contract_call(tmp_path):
    """Trần 62 ngày là của hợp đồng Tracking (`TRAN_NGAY_XUAT`).

    Biết trước thì nói được bằng một câu người đọc hiểu, thay vì nhận
    `khoang-ngay-qua-dai` trở về từ một hệ thống khác — hoặc tệ hơn, gọi một
    lượt mạng chắc chắn hỏng rồi coi đó là "Tracking đang lỗi".
    """
    assert MAX_CONTRACT_DAYS == 62

    narrow = plan_for(tmp_path, write_sales(tmp_path / "hep.xlsx", ROWS[:1], day=SEP))
    assert narrow.day_span == 1 and narrow.fits_one_contract_call

    wide = plan_for(tmp_path, write_sales_over_days(
        tmp_path / "ca-nam.xlsx",
        [(ROWS[0], date(2026, 1, 5)), (ROWS[1], date(2026, 12, 20))],
    ))
    assert wide.date_from == date(2026, 1, 5)
    assert wide.date_to == date(2026, 12, 20)
    assert wide.day_span > MAX_CONTRACT_DAYS
    assert not wide.fits_one_contract_call


def test_an_unreadable_workbook_is_its_own_error_not_a_silent_empty_plan(tmp_path):
    """Nuốt lỗi ở đây thì một sổ hỏng thành một kỳ Pending im lặng."""
    bad = tmp_path / "hong.xlsx"
    bad.write_bytes(b"day khong phai xlsx")
    with pytest.raises(UnreadableSalesWorkbookError):
        plan_for(tmp_path, bad)


# ======================================================================
# 2. Chọn ảnh chụp theo KỲ, không theo "mới nhất toàn cục"
# ======================================================================


def write_daily_min_capture(
    directory: Path, name: str, *, date_from: str, date_to: str,
    captured_at: datetime, codes=("TRK-A",),
) -> Path:
    """Một capture MIN thật, đi qua đúng loader production khi được chọn.

    Trả lời MỌI cặp `(mã, ngày)` trong khoảng — đúng như một lần chụp thật:
    hợp đồng cam kết mỗi cặp đã hỏi nằm ở `records` hoặc ở `errors`, không cặp
    nào bốc hơi. Một fixture chỉ điền ngày đầu sẽ khiến bài kiểm nói về một
    hình dạng dữ liệu mà production không bao giờ nhận được.
    """
    directory.mkdir(parents=True, exist_ok=True)
    first = date.fromisoformat(date_from)
    last = date.fromisoformat(date_to)
    days = [first + timedelta(days=i) for i in range((last - first).days + 1)]
    payload = {
        "capture_id": f"DMIN-{name}",
        "captured_at": captured_at.isoformat(),
        "captured_by": "kiem",
        "source_system_ref": "tracking/api/min-ngay",
        "capture_status": "COMPLETE",
        "data": dmin.contract(
            date_from=date_from, date_to=date_to,
            records=[dmin.record(code, day, min_price=6800)
                     for code in codes for day in days],
        ),
    }
    path = directory / f"{name}.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


@pytest.fixture
def repo(tmp_path: Path) -> Path:
    add_valid_captures(
        tmp_path,
        history_at=CUTOVER,
        catalog_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
    )
    # Danh mục THẬT (có TRK-*) để identity resolve được; bản trong
    # `add_valid_captures` cố ý rỗng.
    write_catalog_capture(tmp_path, CATALOG).replace(
        tmp_path / "data" / "tracking_catalog" / "catalog.json"
    )
    return tmp_path


def test_reopening_an_august_ledger_picks_the_august_capture(repo):
    """Ca trung tâm của bộ này.

    Ảnh chụp tháng 9 MỚI HƠN và hoàn toàn vô dụng cho sổ tháng 8: mọi ngày bán
    nằm ngoài cửa sổ của nó. Chọn theo "mới nhất" thì cả kỳ Pending với lý do
    `SALE_DATE_OUTSIDE_CAPTURE`, và người đọc sẽ đi chụp lại tháng 9 lần nữa.
    """
    kho = repo / "data" / "tracking_daily_min"
    thang_8 = write_daily_min_capture(
        kho, "thang-8", date_from="2026-08-01", date_to="2026-08-31",
        captured_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
    )
    write_daily_min_capture(
        kho, "thang-9", date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
    )

    selected = owner_usability.select_latest_valid_captures(
        repo_root=repo,
        sales=write_sales(repo / "so-thang-8.xlsx", ROWS[:1], day=AUG),
    )
    assert selected.tracking_daily_min == thang_8


def test_a_september_ledger_still_picks_the_september_capture(repo):
    kho = repo / "data" / "tracking_daily_min"
    write_daily_min_capture(
        kho, "thang-8", date_from="2026-08-01", date_to="2026-08-31",
        captured_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
    )
    thang_9 = write_daily_min_capture(
        kho, "thang-9", date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
    )

    selected = owner_usability.select_latest_valid_captures(
        repo_root=repo,
        sales=write_sales(repo / "so-thang-9.xlsx", ROWS[:1], day=SEP),
    )
    assert selected.tracking_daily_min == thang_9


def test_a_capture_that_only_half_covers_the_period_is_not_used(repo):
    """Phủ MỘT PHẦN kỳ vẫn là không trả lời được kỳ ấy. Dùng nó thì nửa đầu
    tháng có giá, nửa sau Pending — và bảng tổng vẫn ra một con số."""
    kho = repo / "data" / "tracking_daily_min"
    write_daily_min_capture(
        kho, "nua-thang", date_from="2026-09-01", date_to="2026-09-10",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
    )
    rows = [ROWS[0]]
    sales = write_sales(repo / "trai-rong.xlsx", rows, day=date(2026, 9, 20))
    selected = owner_usability.select_latest_valid_captures(repo_root=repo, sales=sales)
    assert selected.tracking_daily_min is None


def test_two_captures_of_the_same_period_are_told_apart_by_their_code_set(repo):
    """Cùng kỳ, cùng khoảng ngày, KHÁC tập mã.

    Chuyện này xảy ra một cách rất bình thường: một lần chụp cho sổ của nhân
    viên A, một lần cho sổ của nhân viên B. Cả hai đều "phủ khoảng ngày", nên
    một phép kiểm chỉ nhìn `date_from`/`date_to` sẽ chọn cái MỚI HƠN — và phần
    lớn dòng ra `NOT_IN_CAPTURE`. Đó là một câu trả lời trung thực nhưng nói
    sai vấn đề: người đọc đi tìm hiểu dữ liệu, trong khi việc cần làm là chụp
    lại cho đúng tập mã — và tệ hơn, trong kho ĐANG CÓ một ảnh chụp trả lời
    được.
    """
    kho = repo / "data" / "tracking_daily_min"
    dung_ma = write_daily_min_capture(
        kho, "to-A", date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
        codes=("TRK-A", "TRK-B"),
    )
    write_daily_min_capture(
        kho, "to-B", date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 2, tzinfo=timezone.utc),   # MỚI HƠN
        codes=("TRK-C", "TRK-D"),
    )

    selected = owner_usability.select_latest_valid_captures(
        repo_root=repo,
        sales=write_sales(repo / "so-A.xlsx", ROWS[:2], day=SEP),   # TRK-A, TRK-B
    )
    assert selected.tracking_daily_min == dung_ma


def test_a_capture_missing_one_pair_of_the_period_is_not_used(repo):
    """Thiếu ĐÚNG MỘT cặp cũng là không trả lời được kỳ này.

    Dùng nó thì dòng ấy Pending giữa một bảng có giá — và một dòng thiếu giữa
    một bảng đủ là thứ dễ trôi qua nhất khi đọc.
    """
    kho = repo / "data" / "tracking_daily_min"
    write_daily_min_capture(
        kho, "thieu-mot-ma", date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
        codes=("TRK-A",),
    )
    selected = owner_usability.select_latest_valid_captures(
        repo_root=repo,
        sales=write_sales(repo / "so.xlsx", ROWS[:2], day=SEP),   # cần cả TRK-B
    )
    assert selected.tracking_daily_min is None


def test_a_pair_answered_by_an_error_still_counts_as_answered(repo):
    """"Tracking bảo hôm ấy không có dữ liệu" LÀ một câu trả lời.

    Ảnh chụp mới hơn cũng sẽ nói y như thế, nên coi cặp ấy là "chưa trả lời"
    sẽ loại bỏ một ảnh chụp hoàn toàn hợp lệ và đẩy cả kỳ về Pending vì một
    lý do khác hẳn.
    """
    kho = repo / "data" / "tracking_daily_min"
    directory = kho
    directory.mkdir(parents=True, exist_ok=True)
    payload = {
        "capture_id": "DMIN-co-loi",
        "captured_at": datetime(2026, 10, 1, tzinfo=timezone.utc).isoformat(),
        "captured_by": "kiem",
        "source_system_ref": "tracking/api/min-ngay",
        "capture_status": "COMPLETE",
        "data": dmin.contract(
            date_from="2026-09-03", date_to="2026-09-03",
            records=[dmin.record("TRK-A", "2026-09-03", min_price=6800)],
            errors=[dmin.error("TRK-B", "2026-09-03", "NO_DATA")],
        ),
    }
    path = directory / "co-loi.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")

    selected = owner_usability.select_latest_valid_captures(
        repo_root=repo,
        sales=write_sales(repo / "so.xlsx", ROWS[:2], day=SEP),
    )
    assert selected.tracking_daily_min == path


def test_the_plan_carries_the_exact_pairs_the_ledger_needs(tmp_path):
    """Cặp (mã, ngày) THƯA hơn tích Descartes rất nhiều — và đó là điểm chính.

    Một kỳ 30 ngày với 200 mã có 6.000 ô; sổ thật chỉ chạm vài trăm. Giữ đúng
    những cặp cần trả lời là điều kiện để phép chọn ảnh chụp không đòi hỏi một
    ảnh chụp đầy đủ hơn mức cần thiết.
    """
    plan = plan_for(tmp_path, write_sales_over_days(
        tmp_path / "hai-ngay.xlsx",
        [(ROWS[0], date(2026, 9, 3)), (ROWS[1], date(2026, 9, 20))],
    ))
    assert plan.pairs == (("TRK-A", date(2026, 9, 3)), ("TRK-B", date(2026, 9, 20)))
    assert plan.date_from == date(2026, 9, 3) and plan.date_to == date(2026, 9, 20)


def test_the_plan_splits_a_wide_period_into_contract_sized_windows(tmp_path):
    """Đoạn liền kề, không chồng lấn, phủ đúng khoảng — và mỗi đoạn vừa MỘT
    lượt gọi."""
    plan = plan_for(tmp_path, write_sales_over_days(
        tmp_path / "ca-nam.xlsx",
        [(ROWS[0], date(2026, 1, 5)), (ROWS[1], date(2026, 12, 20))],
    ))
    windows = plan.contract_windows()
    assert len(windows) > 1
    assert all((b - a).days + 1 <= MAX_CONTRACT_DAYS for a, b in windows)
    assert windows[0][0] == plan.date_from
    assert windows[-1][1] == plan.date_to
    for (_, truoc), (sau, _) in zip(windows, windows[1:]):
        assert sau == truoc + timedelta(days=1)


def test_a_narrow_period_is_exactly_one_window(tmp_path):
    plan = plan_for(tmp_path, write_sales(tmp_path / "hep.xlsx", ROWS[:1], day=SEP))
    assert plan.contract_windows() == ((SEP, SEP),)


def test_without_a_workbook_no_daily_min_capture_is_chosen(repo):
    """Không có kỳ thì không có tiêu chí; chọn bừa là đúng lỗi đang sửa."""
    write_daily_min_capture(
        repo / "data" / "tracking_daily_min", "thang-9",
        date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
    )
    selected = owner_usability.select_latest_valid_captures(repo_root=repo)
    assert selected.tracking_daily_min is None


def test_the_chosen_capture_really_answers_the_period(repo):
    """Không chỉ so đường dẫn: nạp lại qua loader production và hỏi thật."""
    kho = repo / "data" / "tracking_daily_min"
    write_daily_min_capture(
        kho, "thang-8", date_from="2026-08-01", date_to="2026-08-31",
        captured_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
    )
    write_daily_min_capture(
        kho, "thang-9", date_from="2026-09-01", date_to="2026-09-30",
        captured_at=datetime(2026, 10, 1, tzinfo=timezone.utc),
    )
    selected = owner_usability.select_latest_valid_captures(
        repo_root=repo,
        sales=write_sales(repo / "so.xlsx", ROWS[:1], day=AUG),
    )
    snapshot = load_daily_min_capture(selected.tracking_daily_min)
    assert snapshot.covers(AUG)
    assert not snapshot.covers(SEP)


# ======================================================================
# 3. Pull-on-run của bản web — có GỌI, có đóng băng, có dọn
# ======================================================================


def fake_tracking(monkeypatch, tmp_path: Path, *, posts: list):
    """Ba capture kia trả về envelope COMPLETE rỗng; chỉ MIN là thật."""
    def envelope(prefix):
        def build(fetch, *, capture_id, captured_by, source_system_ref,
                  captured_at=None, **kw):
            return {
                "capture_id": capture_id,
                "captured_at": (captured_at or datetime.now(timezone.utc)).isoformat(),
                "captured_by": captured_by,
                "source_system_ref": source_system_ref,
                "content_hash": f"hash-{prefix}",
                "capture_status": "COMPLETE",
                **({"rows": CATALOG} if prefix == "cat" else {}),
                **({"entries": {}} if prefix == "inv" else {}),
                **({"nodes": {}} if prefix == "pph" else {}),
            }
        return build

    monkeypatch.setattr(live_pull, "_build_history_capture", envelope("pph"))
    monkeypatch.setattr(live_pull.capture_tracking_catalog, "build_capture", envelope("cat"))
    monkeypatch.setattr(live_pull.capture_inv_map, "build_capture", envelope("inv"))

    def post(body):
        posts.append(body)
        return {
            **dmin.contract(
                date_from=body["date_from"], date_to=body["date_to"],
                records=[dmin.record(code, body["date_from"], min_price=6800)
                         for code in body["product_codes"]],
            ),
            "next_cursor": None,
        }

    return post


def test_the_web_pull_asks_the_contract_once_for_the_whole_period(monkeypatch, tmp_path):
    """Một lượt gọi cho cả kỳ — không phải một lượt cho mỗi dòng bán."""
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)
    sales = write_sales(tmp_path / "so.xlsx", ROWS, day=SEP)

    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test",
        api_key="k", fetch=lambda node: {}, sales=sales, post=post,
    )

    assert len(posts) == 1
    assert posts[0]["product_codes"] == ["TRK-A", "TRK-B", "TRK-C", "TRK-D", "TRK-E"]
    assert (posts[0]["date_from"], posts[0]["date_to"]) == ("2026-09-03", "2026-09-03")
    assert live.tracking_daily_min is not None
    assert live.evidence["daily_min_status"] == "COMPLETE"
    assert live.evidence["daily_min_query_revision"] == dmin.QUERY_REVISION


def test_the_evidence_says_what_the_contract_answered(monkeypatch, tmp_path):
    """Sự cố 2026-09-11: sổ ngày 01–03/09 chạy xong, mọi dòng Tracking `—`, và
    không đâu nói vì sao. `pending_reasons` chỉ ghi `TRACKING_DAILY_MIN_PENDING`
    cho cả `SOURCE_UNAVAILABLE` (Tracking chưa có bản ngày) lẫn `NO_DATA` (có
    bản ngày, mã không có mốc). Bằng chứng của lần chạy phải đếm từng lý do và
    nêu đúng những NGÀY mà mọi mã đều `SOURCE_UNAVAILABLE`."""
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)

    def post_tra_loi_mot_phan(body):
        posts.append(body)
        codes = body["product_codes"]
        d = body["date_from"]
        return {
            **dmin.contract(
                date_from=d, date_to=d,
                records=[dmin.record(codes[0], d, min_price=6800)],
                errors=[dmin.error(c, d, "NO_DATA") for c in codes[1:2]]
                + [dmin.error(c, d, "SOURCE_UNAVAILABLE") for c in codes[2:]],
            ),
            "next_cursor": None,
        }

    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP),
        post=post_tra_loi_mot_phan,
    )
    ev = live.evidence
    assert ev["daily_min_records"] == 1
    assert ev["daily_min_errors"] == 4
    assert ev["daily_min_error_reasons"] == {"NO_DATA": 1, "SOURCE_UNAVAILABLE": 3}
    # 3/5 mã SOURCE_UNAVAILABLE — KHÔNG phải cả ngày: ngày này vẫn có quan sát.
    assert ev["daily_min_unobserved_dates"] == []

    def post_ngay_chua_quan_sat(body):
        d = body["date_from"]
        return {
            **dmin.contract(
                date_from=d, date_to=d, records=[],
                errors=[dmin.error(c, d, "SOURCE_UNAVAILABLE")
                        for c in body["product_codes"]],
            ),
            "next_cursor": None,
        }

    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp2", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, sales=write_sales(tmp_path / "s2.xlsx", ROWS, day=SEP),
        post=post_ngay_chua_quan_sat,
    )
    assert live.evidence["daily_min_records"] == 0
    assert live.evidence["daily_min_error_reasons"] == {"SOURCE_UNAVAILABLE": 5}
    assert live.evidence["daily_min_unobserved_dates"] == [SEP.isoformat()]


def test_the_pulled_capture_is_frozen_for_this_run_and_then_removed(monkeypatch, tmp_path):
    """`S071 §10`: authority thô của Tracking không ở lại trên đĩa máy chủ lâu
    hơn một lần chạy."""
    post = fake_tracking(monkeypatch, tmp_path, posts=[])
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP),
        post=post,
    )
    path = live.tracking_daily_min
    assert path in live.temp_paths and path.is_file()
    assert load_daily_min_capture(path).covers(SEP)

    live.cleanup()
    assert not path.exists()


def test_a_run_without_a_workbook_does_not_call_the_price_contract(monkeypatch, tmp_path):
    """Đường đọc danh mục cho bảng chọn mặt hàng KHÔNG cần giá — và gọi hợp
    đồng giá ở đó là gọi mạng cho một câu hỏi không ai đặt ra."""
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, post=post,
    )
    assert posts == []
    assert live.tracking_daily_min is None
    assert live.evidence["daily_min_skip_reason"] == "NO_SALES_WORKBOOK"


def test_a_failed_price_contract_stops_the_run_instead_of_pricing_nothing(
    monkeypatch, tmp_path
):
    """Từ R1, MIN theo ngày bán LÀ nguồn giá nhập tự động.

    Một báo cáo mà mọi dòng Tracking đều Pending không phải một báo cáo "gần
    đúng" — nó là một báo cáo không có giá vốn, và nó trông y hệt một báo cáo
    có. Thà dừng và nói ra.
    """
    fake_tracking(monkeypatch, tmp_path, posts=[])

    def post(body):
        return {"ok": False, "ly": "nguon-hong"}

    with pytest.raises(live_pull.TrackingUnavailableError) as exc:
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {}, sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP),
            post=post,
        )
    assert exc.value.node == "daily_min"


def test_a_wide_period_is_split_into_windows_and_merged_into_one_capture(
    monkeypatch, tmp_path
):
    """Kỳ rộng hơn 62 ngày phải RA GIÁ, không phải ra một báo cáo rỗng giá.

    Bản trước bỏ qua lượt hỏi giá và lần chạy vẫn tạo báo cáo — đầy đủ hình
    thức, không một giá vốn nào, và trông y hệt một báo cáo bình thường. Nay kỳ
    được chia thành các đoạn ≤ 62 ngày, hỏi từng đoạn, rồi gộp.
    """
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)
    sales = write_sales_over_days(
        tmp_path / "ca-nam.xlsx",
        [(ROWS[0], date(2026, 1, 5)), (ROWS[1], date(2026, 12, 20))],
    )
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, sales=sales, post=post,
    )

    assert len(posts) > 1
    assert all(
        (date.fromisoformat(b["date_to"]) - date.fromisoformat(b["date_from"])).days + 1
        <= MAX_CONTRACT_DAYS
        for b in posts
    )
    # Các đoạn liền kề, không chồng lấn, phủ đúng khoảng của sổ.
    assert posts[0]["date_from"] == "2026-01-05"
    assert posts[-1]["date_to"] == "2026-12-20"
    for truoc, sau in zip(posts, posts[1:]):
        assert date.fromisoformat(sau["date_from"]) == (
            date.fromisoformat(truoc["date_to"]) + timedelta(days=1)
        )

    snapshot = load_daily_min_capture(live.tracking_daily_min)
    assert snapshot.date_from == date(2026, 1, 5)
    assert snapshot.date_to == date(2026, 12, 20)
    assert live.evidence["daily_min_windows"] == len(posts)


def test_windows_of_two_different_states_are_never_merged(monkeypatch, tmp_path):
    """Gộp các đoạn CHỈ hợp lệ khi mọi đoạn cùng `query_revision`.

    Lệch nghĩa là database đã đổi giữa các lượt, và ghép lại thì kỳ báo cáo
    mang giá của hai thời điểm khác nhau. Đây là sự cố thoáng qua (một lượt
    cron chạy đúng lúc), nên nó đi đường "Tracking đang lỗi" — thử lại thật sự
    có tác dụng.
    """
    fake_tracking(monkeypatch, tmp_path, posts=[])
    lan = {"n": 0}

    def post(body):
        lan["n"] += 1
        return {
            **dmin.contract(
                date_from=body["date_from"], date_to=body["date_to"],
                records=[dmin.record(code, body["date_from"], min_price=6800)
                         for code in body["product_codes"]],
                query_revision=f"rev-{lan['n']}",
            ),
            "next_cursor": None,
        }

    sales = write_sales_over_days(
        tmp_path / "ca-nam.xlsx",
        [(ROWS[0], date(2026, 1, 5)), (ROWS[1], date(2026, 12, 20))],
    )
    with pytest.raises(live_pull.TrackingUnavailableError) as exc:
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {}, sales=sales, post=post,
        )
    assert exc.value.reason == "REVISION_CHANGED_MID_CAPTURE"


def test_a_period_beyond_the_window_budget_tells_the_user_to_split_it(
    monkeypatch, tmp_path
):
    """Quá rộng ⇒ nói TÁCH KỲ, không nói "thử lại sau".

    "Thử lại sau" là một lời khuyên sai ở đây: thử bao nhiêu lần cũng thế. Và
    nó KHÔNG được trả về một báo cáo không giá.
    """
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)
    sales = write_sales_over_days(
        tmp_path / "muoi-nam.xlsx",
        [(ROWS[0], date(2020, 1, 5)), (ROWS[1], date(2026, 12, 20))],
    )
    with pytest.raises(live_pull.DailyMinPeriodTooWideError) as exc:
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {}, sales=sales, post=post,
        )
    assert exc.value.windows > live_pull.MAX_CONTRACT_WINDOWS
    assert "tách sổ" in str(exc.value).lower()
    assert posts == []          # không gọi mạng lượt nào


# ======================================================================
# 4. Lần chạy hỏng KHÔNG để lại authority thô trên đĩa
# ======================================================================


def con_lai(tmp_path: Path) -> list[Path]:
    return sorted((tmp_path / "tmp").glob("*.json"))


def test_a_failed_price_contract_leaves_no_capture_behind(monkeypatch, tmp_path):
    """`cleanup()` của bên gọi chỉ chạy khi hàm này TRẢ VỀ một handle.

    Ném ra thì bên gọi không có gì để dọn, và mỗi lần hỏng lại bỏ lại thêm vài
    file capture — đúng thứ `S071 §10` cấm giữ lâu hơn một lần chạy.
    """
    fake_tracking(monkeypatch, tmp_path, posts=[])
    with pytest.raises(live_pull.TrackingUnavailableError):
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {},
            sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP),
            post=lambda body: {"ok": False, "ly": "nguon-hong"},
        )
    assert con_lai(tmp_path) == []


def test_an_exception_inside_planning_leaves_no_capture_behind(monkeypatch, tmp_path):
    """Kể cả một lỗi KHÔNG lường trước từ tận trong kế hoạch hỏi giá.

    Đây là ca mà một khối `except TrackingUnavailableError` hẹp sẽ bỏ lọt —
    nên chỗ dọn dẹp bắt `BaseException`, không bắt riêng một loại.
    """
    fake_tracking(monkeypatch, tmp_path, posts=[])

    def no(*args, **kwargs):
        raise RuntimeError("resolver nổ giữa chừng")

    monkeypatch.setattr(
        "app.modules.pricing.daily_min.planning.plan_daily_min_request_for_workbook", no
    )
    with pytest.raises(RuntimeError):
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {},
            sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP),
            post=lambda body: {},
        )
    assert con_lai(tmp_path) == []


def test_a_period_too_wide_leaves_no_capture_behind(monkeypatch, tmp_path):
    fake_tracking(monkeypatch, tmp_path, posts=[])
    with pytest.raises(live_pull.DailyMinPeriodTooWideError):
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {},
            sales=write_sales_over_days(
                tmp_path / "muoi-nam.xlsx",
                [(ROWS[0], date(2020, 1, 5)), (ROWS[1], date(2026, 12, 20))]),
            post=lambda body: {},
        )
    assert con_lai(tmp_path) == []


def test_a_successful_run_still_keeps_its_captures_until_cleanup(monkeypatch, tmp_path):
    """Đối chứng: đường thành công KHÔNG bị chỗ dọn dẹp mới đụng tới."""
    post = fake_tracking(monkeypatch, tmp_path, posts=[])
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {},
        sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP), post=post,
    )
    assert con_lai(tmp_path) != []
    live.cleanup()
    assert con_lai(tmp_path) == []


# ======================================================================
# 5. Lịch sử `tp/ton` cũ: nguồn LEGACY, không được chặn báo cáo R1
# ======================================================================


def test_a_broken_legacy_history_does_not_block_the_r1_report(monkeypatch, tmp_path):
    """Từ R1, nhánh giá của một mã Tracking đi qua `_daily_min_branch`.

    Lịch sử `tp/ton` chỉ chạy khi caller NÊU RÕ
    `legacy_tracking_history_authority=True`, nên để một sự cố ở nhánh ấy chặn
    cả báo cáo là bắt hôm nay phụ thuộc vào một nguồn hôm nay không dùng.
    Nó vẫn được chụp và vẫn vào bằng chứng — để đối chiếu kết quả sinh trước
    R1 — nhưng vắng mặt thì lần chạy đi tiếp.
    """
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)

    def hong(fetch, *, capture_id, captured_by, source_system_ref,
             captured_at=None, **kw):
        return {
            "capture_id": capture_id,
            "captured_at": (captured_at or datetime.now(timezone.utc)).isoformat(),
            "captured_by": captured_by,
            "source_system_ref": source_system_ref,
            "content_hash": "hash-pph",
            "capture_status": "FAILED",
            "failure_reason": "SOURCE_UNAVAILABLE: 502",
        }

    monkeypatch.setattr(live_pull, "_build_history_capture", hong)
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {},
        sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP), post=post,
    )
    assert live.tracking_capture is None
    assert live.evidence["purchase_price_history_status"] == "FAILED"
    assert "502" in live.evidence["purchase_price_history_failure_reason"]
    # Và giá vẫn được hỏi, vì nguồn giá của R1 không phải nhánh ấy.
    assert len(posts) == 1
    assert live.tracking_daily_min is not None


def test_a_broken_catalog_still_stops_the_run(monkeypatch, tmp_path):
    """Đối chứng: danh mục vẫn REQUIRED. Không có nó thì không resolve được mã
    nào, và mọi dòng Pending vì một lý do identity — không phải vì giá."""
    fake_tracking(monkeypatch, tmp_path, posts=[])

    def hong(fetch, *, capture_id, **kw):
        return {"capture_id": capture_id, "capture_status": "FAILED",
                "failure_reason": "EMPTY_SOURCE_NOT_ASSERTABLE"}

    monkeypatch.setattr(live_pull.capture_tracking_catalog, "build_capture", hong)
    with pytest.raises(live_pull.TrackingUnavailableError) as exc:
        live_pull.pull_live_captures(
            out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
            fetch=lambda node: {},
            sales=write_sales(tmp_path / "s.xlsx", ROWS, day=SEP), post=lambda b: {},
        )
    assert exc.value.node == "catalog"
    assert con_lai(tmp_path) == []


def test_an_unreadable_workbook_does_not_look_like_a_tracking_outage(
    monkeypatch, tmp_path
):
    """Sổ hỏng thì đường nhập sổ ngay sau đây báo lỗi ở nơi người dùng hiểu
    được. Dựng nó thành lỗi Tracking là chỉ sai chỗ cho người đi sửa."""
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)
    bad = tmp_path / "hong.xlsx"
    bad.write_bytes(b"khong phai xlsx")
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, sales=bad, post=post,
    )
    assert posts == []
    assert live.evidence["daily_min_skip_reason"] == "UNREADABLE_SALES_WORKBOOK"


def test_a_ledger_with_no_tracking_lines_is_not_a_tracking_failure(monkeypatch, tmp_path):
    posts: list = []
    post = fake_tracking(monkeypatch, tmp_path, posts=posts)
    rows = [("BH9001", "Tên hàng chưa ai nhận diện", 1, 1_000_000)]
    live = live_pull.pull_live_captures(
        out_dir=tmp_path / "tmp", source_url="https://tracking.test", api_key="k",
        fetch=lambda node: {}, sales=write_sales(tmp_path / "la.xlsx", rows, day=SEP),
        post=post,
    )
    assert posts == []
    assert live.tracking_daily_min is None
    assert live.evidence["daily_min_skip_reason"] == "NO_TRACKING_IDENTITY_LINES"
