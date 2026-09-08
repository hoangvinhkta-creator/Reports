"""R3 — đối soát trên DỮ LIỆU THẬT (hai kỳ golden đã ẩn danh), thành test.

`tools/analysis/r3_reconcile.py` IN ra các con số này để một người đọc. File
này khẳng định những mệnh đề trong số đó mà không được phép hỏng lặng lẽ.

Ranh giới: đây KHÔNG phải Golden Baseline. Nó không so với một file expected
đã commit, và nó không khẳng định doanh thu/lợi nhuận là bao nhiêu — những con
số đó thuộc `tests/test_golden_baseline.py`. Nó khẳng định bốn TÍNH CHẤT của
R3 đúng trên dữ liệu thật:

    §1  nạp lại cùng file, và nạp lại file đảo thứ tự dòng, đều không sinh
        một khoá dòng mới nào
    §2  mọi dòng phụ đo được trên sổ thật đều nhận `FEE`, và không dòng hàng
        thật nào bị nhận nhầm thành dòng phụ
    §3  tổng kỳ == Σ theo nhân viên == Σ theo sheet
    §4  con số trong file .xlsx == con số trên màn hình
"""

from __future__ import annotations

import io
from collections import Counter
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine

import tools.db as history_db
from app.modules.reporting import business_metrics as bm
from app.modules.reporting import line_type as lt
from app.modules.reporting import line_type_config
from app.web import history_store
from tools.analysis import r3_reconcile as rec

@pytest.fixture(scope="module")
def vocabulary():
    return line_type_config.load_vocabulary(rec.CONFIG_DIR / "line_types.yaml")


@pytest.fixture(scope="module", params=rec.PERIODS,
                ids=[spec[0] for spec in rec.PERIODS])
def golden(request):
    filename, period, start, end = request.param
    _result, presented = rec.presented_for(rec.GOLDEN / filename)
    return {"filename": filename, "period": period, "start": start,
            "end": end, "presented": presented}


@pytest.fixture
def loaded(golden, tmp_path):
    """Một database tạm đã nạp kỳ này MỘT lần, cùng dịch vụ đọc nó."""
    engine = create_engine(f"sqlite:///{tmp_path / 'history.db'}")
    history_db.create_all_for_test(engine)
    repository = history_store.SnapshotRepository(engine)
    rec.write_snapshot(repository, golden["presented"], run_id="r1",
                       created_at="2026-09-08T01:00:00", fingerprint="fp-1")
    return {"engine": engine, "repository": repository,
            "service": rec.build_service(engine), **golden}


def current_keys(engine) -> set:
    return {(row.order_key, row.product_key, row.occurrence_index)
            for row in rec._current_keys(engine)}


# --- §1 — nạp lại không làm khoá dòng trôi --------------------------------

def test_reimporting_the_same_book_changes_nothing(loaded):
    before = current_keys(loaded["engine"])
    outcome = rec.write_snapshot(
        loaded["repository"], loaded["presented"], run_id="r2",
        created_at="2026-09-08T02:00:00", fingerprint="fp-1")
    assert outcome.counts["INSERT"] == 0
    assert outcome.counts["SOURCE_CHANGED"] == 0
    assert outcome.counts["SAME"] == len(loaded["presented"])
    assert outcome.ambiguous_bindings == 0
    assert current_keys(loaded["engine"]) == before


def test_reimporting_the_book_with_its_rows_reversed_changes_nothing(loaded):
    """Phép thử trực tiếp của R3 §1 trên sổ THẬT.

    Chỉ VỊ TRÍ dòng đổi; mọi giá trị nghiệp vụ giữ nguyên tuyệt đối. Trước R3
    phép biến đổi này làm `occurrence_index` của mọi nhóm nhiều dòng đảo lại —
    và một quyết định của Owner đi theo chỉ số ấy sẽ trôi sang dòng khác.
    """
    before = current_keys(loaded["engine"])
    outcome = rec.write_snapshot(
        loaded["repository"], loaded["presented"], run_id="r3",
        created_at="2026-09-08T03:00:00", fingerprint="fp-2",
        source_lines=rec.reversed_rows(loaded["presented"]))
    assert outcome.counts["INSERT"] == 0
    assert outcome.counts["SOURCE_CHANGED"] == 0
    assert outcome.ambiguous_bindings == 0
    assert current_keys(loaded["engine"]) == before


def test_the_golden_books_contain_no_order_with_a_repeated_item(golden):
    """GIỚI HẠN của dữ liệu thật khả dụng, nói ra chứ không giấu.

    Hai kỳ golden đã ẩn danh KHÔNG có đơn nào chứa hai dòng cùng một mặt hàng
    (nhóm lớn nhất = 1). Nghĩa là bài đảo-thứ-tự ở trên, chạy trên chính hai
    kỳ này, KHÔNG chạm tới ca dễ hỏng nhất — nó chỉ chứng minh phép gắn không
    làm hỏng ca thường.

    Ca dễ hỏng được dựng tường minh ở bài ngay dưới, TRÊN NỀN chính sổ thật
    này. Nếu một ngày dữ liệu golden có nhóm nhiều dòng thật, bài này sẽ đỏ và
    đó là tín hiệu để bỏ nó đi, không phải để nới lỏng nó.
    """
    from app.history import extraction
    groups = Counter(
        (item.key.order_key, item.key.product_key)
        for item in extraction.build_source_lines(golden["presented"]))
    assert max(groups.values()) == 1


def test_a_repeated_item_built_on_the_real_book_survives_a_reorder(
    loaded, tmp_path
):
    """Ca dễ hỏng nhất, dựng TRÊN NỀN sổ thật, đi hết qua database.

    Lấy một dòng phí CÓ THẬT của kỳ, nhân đôi nó vào cùng đơn với một số tiền
    khác — đúng hình dạng mà `extraction.build_source_lines` mô tả là có trong
    sổ gốc ("một đơn chứa hai dòng cùng tên hàng"). Gắn một giá nhập tay lên
    dòng thứ nhất, rồi nạp lại sổ với HAI DÒNG ĐÓ ĐẢO CHỖ.

    Trước R3, dòng đứng trên nhận `occurrence_index = 1` và giá nhập trôi
    sang dòng kia. Bài này khẳng định nó không còn trôi.
    """
    from dataclasses import replace as dc_replace
    from decimal import Decimal as D

    from sqlalchemy import insert, select

    from app.history import extraction, keys
    from tools.db import schema

    source = extraction.build_source_lines(loaded["presented"])
    fee = next((item for item in source
                if "chi phí" in (item.product_raw or "").casefold()), None)
    assert fee is not None, "kỳ này không có dòng phí nào để nhân đôi"

    def clone(item, *, row, occurrence, sell):
        values = dict(sale_date=item.sale_date, product_raw=item.product_raw,
                      quantity=item.quantity, sell_price=D(sell),
                      discount=item.discount, total_sales_raw=D(sell),
                      delivery_cost=item.delivery_cost, imei=item.imei,
                      note_raw=item.note_raw, employee_raw=item.employee_raw,
                      source_profit=item.source_profit)
        ordered = tuple(values[name] for name in keys.FINGERPRINT_FIELDS)
        return dc_replace(
            item, key=type(item.key)(item.key.order_key, item.key.product_key,
                                     occurrence),
            source_row=row, fingerprint=keys.line_fingerprint(ordered), **values)

    twins = [clone(fee, row=10_001, occurrence=1, sell="100000"),
             clone(fee, row=10_002, occurrence=2, sell="200000")]
    repository = loaded["repository"]
    rec.write_snapshot(repository, loaded["presented"], run_id="twin-1",
                       created_at="2026-09-08T04:00:00", fingerprint="fp-twin",
                       source_lines=twins)

    engine = loaded["engine"]
    with engine.begin() as connection:
        connection.execute(insert(schema.kpi_purchase_price_override).values(
            order_key=fee.key.order_key, product_key=fee.key.product_key,
            occurrence_index=1, origin=schema.ORIGIN_PIPELINE,
            purchase_price=D("70000"), provenance="MANUAL",
            auto_price_at_entry=None, entered_at="2026-09-08T04:30:00",
            entered_by="owner-web", reason="Hàng ngoài bảng giá"))

    def sell_price_of(occurrence):
        with engine.connect() as connection:
            return connection.execute(
                select(schema.order_line_source_version.c.sell_price)
                .select_from(schema.order_line_current.join(
                    schema.order_line_source_version,
                    schema.order_line_source_version.c.id
                    == schema.order_line_current.c.current_source_version_id))
                .where(schema.order_line_current.c.order_key == fee.key.order_key)
                .where(schema.order_line_current.c.product_key == fee.key.product_key)
                .where(schema.order_line_current.c.occurrence_index == occurrence)
            ).scalar()

    assert sell_price_of(1) == D("100000")

    swapped = [clone(fee, row=10_001, occurrence=1, sell="200000"),
               clone(fee, row=10_002, occurrence=2, sell="100000")]
    outcome = rec.write_snapshot(
        repository, loaded["presented"], run_id="twin-2",
        created_at="2026-09-08T05:00:00", fingerprint="fp-twin-2",
        source_lines=swapped)

    # Khoá số 1 VẪN là dòng 100.000 — giá nhập Owner gõ không trôi đi đâu.
    assert sell_price_of(1) == D("100000")
    assert sell_price_of(2) == D("200000")
    assert outcome.counts["SOURCE_CHANGED"] == 0
    assert outcome.ambiguous_bindings == 0


# --- §2 — loại dòng trên sổ thật ------------------------------------------

def test_every_supplementary_label_of_the_real_book_is_classified_as_a_fee(
    golden, vocabulary
):
    """Các nhãn dòng phụ ĐO ĐƯỢC trên hai kỳ này (`Chi phí vận chuyển`,
    `Chi phí lắp đặt`, `Chênh VAT`, `Phụ Phí`…) đều phải ra `FEE`.

    Đây là điều kiện để `TASK-105B-Q3` thật sự mở khoá: một nhãn bị bỏ sót là
    một dòng vĩnh viễn không có giá nhập, và cả kỳ vĩnh viễn không CHÍNH THỨC.
    """
    missed = []
    for view in golden["presented"]:
        product = (view.line.product_raw or "").casefold()
        looks_supplementary = any(
            token in product for token in
            ("chi phí", "phụ phí", "chênh vat", "công lắp đặt"))
        if not looks_supplementary:
            continue
        kind = lt.classify(
            vocabulary=vocabulary, order_key=view.line.order_id,
            product_raw=view.line.product_raw, sell_price=view.line.sell_price,
            quantity=view.line.quantity)
        # Một dòng phụ nằm trong chứng từ chưa định nghĩa vẫn đúng khi ra
        # `UNDECIDED_DOCUMENT` — chứng từ thắng dòng (`classify` bước 1).
        if kind not in (lt.TYPE_FEE, lt.TYPE_UNDECIDED_DOCUMENT):
            missed.append((view.line.order_id, view.line.product_raw, kind))
    assert missed == []


def test_no_real_product_line_is_mistaken_for_a_fee(golden, vocabulary, loaded):
    """Chiều ngược lại: một mặt hàng thật bị nhận nhầm thành dòng phụ sẽ
    lặng lẽ nhận giá nhập `0` và thổi phồng lợi nhuận."""
    wrong = []
    for line, detail in zip(loaded["service"].period(
            date_from=loaded["start"], date_to=loaded["end"]).lines,
            loaded["service"].period(
                date_from=loaded["start"], date_to=loaded["end"]).details):
        if line.line_type != lt.TYPE_FEE:
            continue
        product = (detail["product_raw"] or "").casefold()
        if not any(token in product for token in
                   ("phí", "cước", "chênh vat", "công lắp đặt")):
            wrong.append(detail["product_raw"])
    assert wrong == []


def test_the_fee_lines_are_the_ones_carrying_the_policy_price(loaded):
    """Số dòng nhận giá `0` theo chính sách đúng bằng số dòng `FEE`/`DISCOUNT`
    KHÔNG có nguồn giá thật — không nhiều hơn một dòng nào."""
    data = loaded["service"].period(
        date_from=loaded["start"], date_to=loaded["end"])
    policy = [line for line in data.lines
              if line.purchase_provenance == bm.PROVENANCE_POLICY_ZERO]
    assert policy, "kỳ này không có dòng phụ nào — bài test mất ý nghĩa"
    assert all(line.line_type in lt.POLICY_ZERO_TYPES for line in policy)
    assert all(line.purchase_price == Decimal(0) for line in policy)
    # Và không dòng hàng thật nào rơi vào đó.
    assert not [line for line in data.lines
                if line.line_type == lt.TYPE_SALE
                and line.purchase_provenance == bm.PROVENANCE_POLICY_ZERO]


def test_a_missing_price_on_a_real_product_line_is_still_missing(loaded):
    """`OD-105B-01` §3 câu 2, đo trên sổ thật: hàng bán thiếu giá KHÔNG hoá 0."""
    data = loaded["service"].period(
        date_from=loaded["start"], date_to=loaded["end"])
    unpriced = [line for line in data.lines
                if line.line_type == lt.TYPE_SALE and line.purchase_price is None]
    assert unpriced, "kỳ này không còn dòng thiếu giá — bài test mất ý nghĩa"
    assert all(line.kpi_profit is None for line in unpriced)


# --- §3 — một effective data ----------------------------------------------

def test_the_period_the_employees_and_the_sheets_all_add_up(loaded):
    service = loaded["service"]
    data = service.period(date_from=loaded["start"], date_to=loaded["end"])
    by_employee = bm.group_by_employee(data.lines)
    by_sheet = [data.for_sheet(sheet).totals for sheet in service.sheets(data)]

    assert sum(t.lines for _n, _g, t in by_employee) == data.totals.lines
    assert sum(t.lines for t in by_sheet) == data.totals.lines
    assert rec._sum(t.sales_revenue for _n, _g, t in by_employee) == \
        data.totals.sales_revenue
    assert rec._sum(t.sales_revenue for t in by_sheet) == data.totals.sales_revenue
    assert rec._sum(t.kpi_profit for _n, _g, t in by_employee) == \
        data.totals.kpi_profit


# --- §4 — file xuất nói cùng câu với màn hình -----------------------------

def test_the_workbook_matches_the_screen(loaded):
    service = loaded["service"]
    data = service.period(date_from=loaded["start"], date_to=loaded["end"],
                          period=loaded["period"])
    stream = io.BytesIO()
    service.export_workbook(data=data, period_label="kỳ").save(stream)
    stream.seek(0)
    book = load_workbook(stream)
    exported = rec._exported_totals(book)
    assert exported["lines"] == data.totals.lines
    assert Decimal(str(exported["kpi_profit"])) == data.totals.kpi_profit
    # Mỗi dòng đúng MỘT lần trong các sheet dòng — phân hoạch, không phải lọc.
    line_sheets = [name for name in book.sheetnames
                   if name not in ("Tổng kỳ", "Theo nhân viên")]
    assert sum(book[name].max_row - 2 for name in line_sheets) == data.totals.lines


def test_the_workbook_never_writes_zero_for_a_missing_price(loaded):
    """Ô Giá nhập KPI trống nghĩa là CHƯA CÓ GIÁ — không bao giờ là `0`."""
    service = loaded["service"]
    data = service.period(date_from=loaded["start"], date_to=loaded["end"])
    stream = io.BytesIO()
    service.export_workbook(data=data, period_label="kỳ").save(stream)
    stream.seek(0)
    book = load_workbook(stream)
    written = []
    for name in book.sheetnames:
        if name in ("Tổng kỳ", "Theo nhân viên"):
            continue
        sheet = book[name]
        headers = [cell.value for cell in sheet[1]]
        column = headers.index("Giá nhập KPI") + 1
        kind = headers.index("Loại dòng") + 1
        for row in range(2, sheet.max_row):  # bỏ dòng TỔNG cuối
            written.append((sheet.cell(row=row, column=kind).value,
                            sheet.cell(row=row, column=column).value))
    zeros = [kind for kind, value in written if value == 0]
    assert zeros, "không có dòng nào giá 0 — bài test mất ý nghĩa"
    # Mọi số `0` đã ghi ra đều thuộc một loại dòng CÓ THẨM QUYỀN cho số đó.
    allowed = {lt.LINE_TYPE_LABELS[name] for name in lt.POLICY_ZERO_TYPES}
    assert set(zeros) <= allowed
