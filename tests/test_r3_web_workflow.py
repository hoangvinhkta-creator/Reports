"""R3 — bốn luồng của Owner, kiểm qua ỨNG DỤNG WEB THẬT.

`tests/test_r3_line_types.py` và `tests/test_r3_export_and_period_close.py`
kiểm ngữ nghĩa ở tầng domain/dịch vụ. File này hỏi câu Owner thật sự hỏi:

    tôi bấm nút trên trang → nó có làm đúng cái nó hứa không, và có còn đúng
    sau khi tôi mở lại trang không?

Bốn luồng: tải file Excel · chốt kỳ · thao tác bị TỪ CHỐI sau khi chốt · mở
lại kỳ. Toàn bộ dữ liệu là tổng hợp.
"""

from __future__ import annotations

import io
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine

import tools.db as history_db
from app.modules.exporting import business_export
from app.modules.reporting import business_metrics as bm
from app.web import business_service, business_store, history_store
from app.web import server as web_server
from tests.test_employee_workspace_ux import (
    SEPTEMBER, TODAY, body, line, metric, metrics, persist,
)
from tools.tracking import live_pull

PERIOD_QS = "ky=2026-09"


@pytest.fixture
def engine():
    engine = create_engine("sqlite://")
    history_db.create_all_for_test(engine)
    return engine


@pytest.fixture
def repository(engine):
    return history_store.SnapshotRepository(engine)


@pytest.fixture
def store(engine):
    return business_store.BusinessDecisionStore(engine)


@pytest.fixture
def service(engine, store):
    return business_service.BusinessReportService(engine=engine, store=store)


@pytest.fixture
def client(engine, monkeypatch, tmp_path):
    monkeypatch.setattr(web_server, "select_latest_valid_captures", lambda: None)
    monkeypatch.setattr(live_pull, "is_configured", lambda env=None: False)
    monkeypatch.setattr(web_server, "_today", lambda: TODAY)
    monkeypatch.setattr(
        web_server.identity_gateway, "DEFAULT_LOG_PATH",
        tmp_path / "identity" / "mappings.jsonl")
    monkeypatch.setattr(
        web_server.identity_gateway, "DEFAULT_INDEX_PATH",
        tmp_path / "identity" / "index.json")
    application = web_server.create_app(
        db_path=tmp_path / "runs.db",
        history=history_store.LegacyRepository(engine),
        snapshots=history_store.SnapshotRepository(engine))
    application.testing = True
    return application.test_client()


def one_order():
    return [line("BH72707", "43F6000", day=5, sell="8000000",
                 kpi_purchase="5000000", kpi_profit="3000000")]


def keys_of(service, order_key, product):
    for detail in service.period(**SEPTEMBER).details:
        if detail["order_key"] == order_key and detail["product_raw"] == product:
            return {"order_key": detail["order_key"],
                    "product_key": detail["product_key"],
                    "occurrence_index": str(detail["occurrence_index"])}
    raise AssertionError(f"không tìm thấy dòng {order_key}/{product}")


# --------------------------------------------------------------------------
# §4 — tải file Excel
# --------------------------------------------------------------------------

class TestDownloadingTheWorkbook:
    def test_the_download_is_an_xlsx_named_after_the_period(
        self, repository, client
    ):
        persist(repository, one_order())
        response = client.get(f"/kinh-doanh/xuat-excel?{PERIOD_QS}")
        assert response.status_code == 200
        assert response.mimetype == web_server.XLSX_MIMETYPE
        assert "bao-cao-2026-09-toan-ky.xlsx" in \
            response.headers["Content-Disposition"]

    def test_the_file_carries_the_price_the_owner_just_typed(
        self, repository, service, client
    ):
        """Cùng một mệnh đề như bài domain, nhưng qua ĐÚNG hai route thật:
        Owner lưu giá ở một trang, rồi tải file ở một trang khác."""
        persist(repository, one_order())
        keys = keys_of(service, "BH72707", "43F6000")
        saved = client.post("/kinh-doanh/gia-nhap", data={
            "ky": "2026-09", "gia_nhap": "4.000.000",
            "ly_do": "Giá nhập theo hoá đơn", **keys})
        assert saved.status_code == 302

        response = client.get(f"/kinh-doanh/xuat-excel?{PERIOD_QS}")
        book = load_workbook(io.BytesIO(response.data))
        sheet = book["Nội thành"]
        headers = [cell.value for cell in sheet[1]]
        price = sheet.cell(row=2, column=headers.index("Giá nhập KPI") + 1).value
        source = sheet.cell(
            row=2, column=headers.index("Nguồn giá nhập") + 1).value
        assert price == 4000000
        assert source == business_export.PROVENANCE_LABELS[
            bm.PROVENANCE_MANUAL_OVERRIDE]

    def test_a_single_employee_can_be_downloaded_on_its_own(
        self, repository, client
    ):
        persist(repository, one_order())
        response = client.get(
            f"/kinh-doanh/xuat-excel?{PERIOD_QS}&nhan-vien=Vinh")
        assert response.status_code == 200
        assert "Vinh" in response.headers["Content-Disposition"]

    def test_an_unknown_sheet_is_a_404_not_an_empty_workbook(
        self, repository, client
    ):
        """Một file rỗng trông như "kỳ này không có gì" — một câu sai."""
        persist(repository, one_order())
        assert client.get(
            f"/kinh-doanh/xuat-excel?{PERIOD_QS}&nhom=khong-co-that"
        ).status_code == 404


# --------------------------------------------------------------------------
# §5 — chốt kỳ, và cái khoá đi kèm
# --------------------------------------------------------------------------

class TestClosingThePeriodThroughTheWeb:
    def test_the_page_shows_the_same_numbers_as_the_report(
        self, repository, service, client
    ):
        persist(repository, one_order())
        html = body(client, f"/kinh-doanh/chot-ky?{PERIOD_QS}")
        totals = service.period(**SEPTEMBER).totals
        assert metric(html, "close-state") == "CHƯA CHỐT"
        assert metric(html, "close-lines") == str(totals.lines)

    def test_closing_then_saving_a_price_is_refused(
        self, repository, service, client
    ):
        """Bài trung tâm của §5: cái chốt phải TỪ CHỐI, không chỉ cảnh báo."""
        persist(repository, one_order())
        keys = keys_of(service, "BH72707", "43F6000")

        closed = client.post("/kinh-doanh/chot-ky", data={
            "ky": "2026-09", "ghi_chu": "Đã duyệt tháng 9"})
        assert closed.status_code == 302

        refused = client.post("/kinh-doanh/gia-nhap", data={
            "ky": "2026-09", "gia_nhap": "1.000.000",
            "ly_do": "sửa sau khi chốt", **keys})
        assert refused.status_code == 409
        assert "đã được chốt" in refused.get_data(as_text=True)
        # Và con số KHÔNG đổi — từ chối nghĩa là không ghi gì cả.
        assert service.period(**SEPTEMBER).lines[0].purchase_price == \
            Decimal("5000000")

    def test_a_line_of_another_month_is_still_editable(
        self, repository, service, client
    ):
        """Chốt tháng 9 KHÔNG được khoá tháng 10."""
        persist(repository, [
            line("BH72707", "43F6000", day=5, month=9, row=6),
            line("BH72999", "43F6000", day=5, month=10, row=7),
        ])
        client.post("/kinh-doanh/chot-ky", data={"ky": "2026-09"})
        keys = None
        for detail in service.period(date_from=None, date_to=None).details:
            if detail["order_key"] == "BH72999":
                keys = {"order_key": detail["order_key"],
                        "product_key": detail["product_key"],
                        "occurrence_index": str(detail["occurrence_index"])}
        assert keys is not None
        saved = client.post("/kinh-doanh/gia-nhap", data={
            "ky": "tat-ca", "gia_nhap": "4.000.000",
            "ly_do": "tháng 10 vẫn mở", **keys})
        assert saved.status_code == 302

    def test_reopening_needs_a_reason_and_then_unlocks(
        self, repository, service, client
    ):
        persist(repository, one_order())
        keys = keys_of(service, "BH72707", "43F6000")
        client.post("/kinh-doanh/chot-ky", data={"ky": "2026-09"})

        blank = client.post("/kinh-doanh/chot-ky", data={
            "ky": "2026-09", "hanh-dong": "mo-lai", "ly_do": "   "})
        assert blank.status_code == 302
        html = body(client, f"/kinh-doanh/chot-ky?{PERIOD_QS}")
        assert "ĐÃ CHỐT" in metric(html, "close-state")  # vẫn còn khoá

        client.post("/kinh-doanh/chot-ky", data={
            "ky": "2026-09", "hanh-dong": "mo-lai",
            "ly_do": "Sót một hoá đơn của ngày 30"})
        saved = client.post("/kinh-doanh/gia-nhap", data={
            "ky": "2026-09", "gia_nhap": "4.000.000",
            "ly_do": "bổ sung sau khi mở lại", **keys})
        assert saved.status_code == 302

    def test_the_close_survives_a_brand_new_application(
        self, engine, repository, client, monkeypatch, tmp_path
    ):
        """Chốt kỳ phải sống qua redeploy — nó nằm trong database, không RAM."""
        persist(repository, one_order())
        client.post("/kinh-doanh/chot-ky", data={
            "ky": "2026-09", "ghi_chu": "Đã duyệt"})

        fresh = web_server.create_app(
            db_path=tmp_path / "runs2.db",
            history=history_store.LegacyRepository(engine),
            snapshots=history_store.SnapshotRepository(engine))
        fresh.testing = True
        html = body(fresh.test_client(), f"/kinh-doanh/chot-ky?{PERIOD_QS}")
        assert "ĐÃ CHỐT" in metric(html, "close-state")
        assert "Đã duyệt" in html

    def test_the_all_data_view_cannot_be_closed(self, repository, client):
        persist(repository, one_order())
        assert client.post("/kinh-doanh/chot-ky",
                           data={"ky": "tat-ca"}).status_code == 404


# --------------------------------------------------------------------------
# §2/§1 — hai hàng đợi ngoại lệ mới có mặt trên bảng kê
# --------------------------------------------------------------------------

class TestTheNewExceptionQueuesAreReachable:
    def test_the_policy_zero_queue_lists_the_fee_lines(self, repository, client):
        persist(repository, [
            line("BH72707", "43F6000", day=5, row=6),
            line("BH72707", "Chi phí vận chuyển", day=5, row=7, sell="0",
                 kpi_purchase=None, kpi_profit=None,
                 reasons=("Missing.PurchasePrice",)),
        ])
        html = body(
            client,
            f"/kinh-doanh/gia-nhap?{PERIOD_QS}&loc=gia-theo-chinh-sach")
        assert "Chi phí vận chuyển" in html
        assert "43F6000" not in html
        # Và nó hiện đúng NGUỒN của con số 0 ấy — không phải "Tự động".
        assert "Chính sách (dòng phụ)" in metrics(html, "provenance")[0]

    def test_the_undecided_document_queue_lists_the_btl_order(
        self, repository, client
    ):
        persist(repository, [
            line("BH72707", "43F6000", day=5, row=6),
            line("BTL00301", "43F6000", day=5, row=7, kpi_purchase=None,
                 kpi_profit=None, reasons=("Missing.PurchasePrice",)),
        ])
        html = body(client,
                    f"/kinh-doanh/gia-nhap?{PERIOD_QS}&loc=loai-chua-ro")
        assert "BTL00301" in html
        assert "BH72707" not in html
